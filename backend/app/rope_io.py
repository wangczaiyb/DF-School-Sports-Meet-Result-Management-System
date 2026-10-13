"""跳绳报名表识别、抽签赛程及可打印的登记/汇总表。"""
from __future__ import annotations

import csv
import math
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from random import SystemRandom
from urllib.parse import quote
from zipfile import ZipFile

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .database import get_db
from .dependencies import business_admin, manager_only
from .models import (CollegeTeam, CollegeTeamMember, Event, EventRound, EventSchedule,
                     EventAdminPermission, Registration, Result, RopeLineupMember, Role, SportsMeet, StudentRoster, User)
from .rope_rules import ROPE_RULES, rope_lineup_status
from .services import audit, require_event_permission, result_dict
from .college_policy import configured_colleges

router = APIRouter()
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_TITLES = {"ROPE_SINGLE": "计时计数项目", "ROPE_SYNC": "集体单摇跳", "ROPE_FIG8": "8字绕绳跳"}
ALIASES = {
    "student_no": {"学号", "学生学号", "学生编号", "学生号"},
    "name": {"姓名", "学生姓名", "运动员姓名", "参赛者姓名"},
    "gender": {"性别", "男女"},
    "college": {"学院", "院系", "所属学院", "参赛学院", "代表队", "单位", "系别"},
    "event_code": {"项目", "参赛项目", "比赛项目", "报名项目", "小项目"},
    "role": {"角色", "参赛角色", "分工", "队员角色"},
    "major": {"专业"}, "class_name": {"班级", "行政班", "班级名称"},
}


@router.get("/admin/rope/meets")
def available_rope_meets(actor: User = Depends(manager_only), db: Session = Depends(get_db)):
    query = select(SportsMeet).join(Event, Event.meet_id == SportsMeet.id).where(Event.code.in_(ROPE_RULES))
    if actor.role not in {Role.DEVELOPER, Role.SPORTS_TEACHER}:
        query = query.join(EventAdminPermission, EventAdminPermission.event_id == Event.id).where(
            EventAdminPermission.user_id == actor.id,
            (EventAdminPermission.can_schedule | EventAdminPermission.can_score | EventAdminPermission.can_review | EventAdminPermission.can_export))
    return list(db.scalars(query.distinct().order_by(SportsMeet.start_date.desc())))


@router.get("/admin/rope/status")
def admin_rope_status(meet_id: int, actor: User = Depends(manager_only), db: Session = Depends(get_db)):
    from .api import rope_status
    status = rope_status(_actor=actor, meet_id=meet_id, db=db)
    if not status["events"] and actor.role not in {Role.DEVELOPER, Role.SPORTS_TEACHER}:
        raise HTTPException(403, "没有该运动会跳绳项目的查看权限")
    return status


def clean(value) -> str:
    return re.sub(r"[\s*：:（）()\"“”]", "", str(value or "")).strip()


def event_code(value) -> str:
    value = clean(value)
    if value in ROPE_RULES:
        return value
    if any(word in value for word in ("前单跳", "计时计数", "单人跳", "一分钟跳", "1分钟跳")):
        return "ROPE_SINGLE"
    if "集体单摇" in value or "集体跳" in value:
        return "ROPE_SYNC"
    if "8" in value and any(word in value for word in ("字", "绕绳", "环绕")):
        return "ROPE_FIG8"
    return ""


def event_codes(value) -> list[str]:
    """一格填写多个项目时保留每个报名项目。"""
    return list(dict.fromkeys(code for part in re.split(r"[,，、;；/|\n]+", str(value)) if (code := event_code(part))))


def cell_text(cell) -> str:
    value = cell.value
    if cell.data_type == "f":
        raise ValueError("个人信息不能使用公式，请粘贴为文本或数值")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not value.is_integer():
            return str(value)
        if re.fullmatch(r"0+", cell.number_format or ""):
            return str(int(value)).zfill(len(cell.number_format))
        return str(int(value))
    return str(value).strip() if value is not None else ""


def read_registration(file: UploadFile, college: str, default_event: str, last_two_turners: bool) -> dict:
    filename = file.filename or ""
    content = file.file.read(8 * 1024 * 1024 + 1)
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(422, "报名表不能超过8MB")
    workbook = None
    try:
        if filename.lower().endswith(".csv"):
            try:
                source = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                source = content.decode("gb18030")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "报名表"
            for row in csv.reader(StringIO(source)):
                sheet.append(row)
        elif filename.lower().endswith(".xlsx"):
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                    raise HTTPException(422, "Excel解压后过大，请拆分报名表")
            workbook = load_workbook(BytesIO(content), data_only=False)
        else:
            raise HTTPException(422, "请上传.xlsx或.csv文件；旧版.xls请先另存为.xlsx")
        records, errors, warnings, sheets = [], [], [], []
        for sheet in workbook:
            if sheet.max_row > 10000 or sheet.max_column > 100:
                errors.append(f"{sheet.title}超出10000行/100列限制，请拆分表格")
                continue
            columns = {}
            context_college = college.strip() or (sheet.title if sheet.title.endswith("学院") else "")
            context_event = event_code(sheet.title) or default_event
            sheet_count = 0
            # 兼容纵向合并的学院和项目列，保留原表行号用于报错。
            for merged in list(sheet.merged_cells.ranges):
                if merged.min_col == merged.max_col and merged.max_row - merged.min_row < 10000:
                    value = sheet.cell(merged.min_row, merged.min_col).value
                    sheet.unmerge_cells(str(merged))
                    for row_no in range(merged.min_row, merged.max_row + 1):
                        sheet.cell(row_no, merged.min_col, value)
            for row_no, cells in enumerate(sheet.iter_rows(), 1):
                values = [str(cell.value).strip() if cell.value is not None else "" for cell in cells]
                found = {field: index for index, value in enumerate(values)
                         for field, aliases in ALIASES.items() if clean(value) in aliases}
                if "student_no" in found and "name" in found:
                    found.update({f"project_{code}": index for index, value in enumerate(values)
                                  if (code := event_code(value))})
                    columns = found
                    continue
                first = next((value for value in values if value), "")
                metadata = re.match(r"^(代表队|学院|院系)\s*[:：]\s*(.*)$", first)
                if metadata:
                    context_college = metadata.group(2).strip() or next((v for v in values[1:] if v), college.strip())
                    continue
                title_event = event_code(first)
                if title_event and sum(bool(value) for value in values) == 1:
                    context_event = title_event
                    continue
                if not columns:
                    continue
                no_value = values[columns["student_no"]]
                name_value = values[columns["name"]]
                if not no_value and not name_value:
                    continue
                location = f"{sheet.title}第{row_no}行"
                try:
                    record = {field: cell_text(cells[index]) for field, index in columns.items()}
                except ValueError as exc:
                    errors.append(f"{location}：{exc}")
                    continue
                no = record.get("student_no", "")
                if not 6 <= len(no) <= 32 or not re.fullmatch(r"[A-Za-z0-9_-]+", no):
                    errors.append(f"{location}：学号须为6至32位字母/数字，请将学号列设为文本")
                    continue
                if isinstance(cells[columns["student_no"]].value, (int, float)) and len(no) > 15:
                    errors.append(f"{location}：超过15位的数字学号可能丢失精度，请使用文本学号")
                record["gender"] = {"男生": "男", "男性": "男", "M": "男", "MALE": "男",
                                    "女生": "女", "女性": "女", "F": "女", "FEMALE": "女"}.get(
                                        record.get("gender", "").upper(), record.get("gender", ""))
                record["college"] = record.get("college") or context_college
                raw_event = record.get("event_code", "")
                if raw_event and not event_code(raw_event):
                    errors.append(f"{location}：无法识别项目“{raw_event}”")
                codes = event_codes(raw_event) if raw_event else []
                project_columns = {field.removeprefix("project_"): value for field, value in record.items() if field.startswith("project_")}
                if project_columns:
                    codes = [code for code, value in project_columns.items() if value and clean(value).upper() not in {"0", "否", "未报名", "×", "X", "FALSE"}]
                elif not codes:
                    codes = [context_event] if context_event else []
                raw_role = record.get("role", "")
                record["role"] = {"跳绳者": "JUMPER", "跳绳": "JUMPER", "队员": "JUMPER", "JUMPER": "JUMPER",
                                  "摇绳者": "TURNER", "摇绳": "TURNER", "TURNER": "TURNER"}.get(raw_role.upper(), "")
                if raw_role and not record["role"]:
                    errors.append(f"{location}：角色应为跳绳者或摇绳者")
                if not record["name"] or len(record["name"]) > 64 or record["gender"] not in {"男", "女"} or not record["college"]:
                    errors.append(f"{location}：姓名、性别（男/女）、学院必须完整，请补充表格或填写默认学院")
                if any(len(record.get(field, "")) > 128 for field in ("college", "major", "class_name")):
                    errors.append(f"{location}：学院、专业或班级不能超过128字")
                record.update(sheet=sheet.title, row=row_no)
                for code in codes or [""]:
                    records.append({**record, "event_code": code})
                sheet_count += 1
            sheets.append({"sheet": sheet.title, "count": sheet_count})
        raw_entries = len(records)
        unique, identities, lineups = {}, {}, defaultdict(list)
        for record in records:
            no = record["student_no"]
            identity = (record["name"], record["gender"], record["college"])
            if no in identities and identities[no] != identity:
                errors.append(f"学号{no}的姓名、性别或学院在不同位置不一致")
            identities[no] = identity
            key = (no, record["event_code"], record["college"])
            if key in unique and unique[key]["role"] != record["role"]:
                errors.append(f"学号{no}在同一项目的角色不一致")
            unique.setdefault(key, record)
        records = list(unique.values())
        for record in records:
            if record["event_code"]:
                lineups[(record["college"], record["event_code"])].append(record)
        for (academy, code), members in lineups.items():
            missing_roles = [record for record in members if not record["role"]]
            if missing_roles and ROPE_RULES[code]["turner_count"]:
                if last_two_turners and len(missing_roles) == len(members) and len(members) == ROPE_RULES[code]["lineup_size"]:
                    for record in members[-2:]:
                        record["role"] = "TURNER"
                else:
                    warnings.append(f"{academy}的{SHEET_TITLES[code]}未标完整角色，导入后请配置摇绳者")
            for record in members:
                record["role"] = record["role"] or "JUMPER"
        if not records:
            errors.append("没有识别到报名人员。空白成绩模板需先填写学号、姓名、性别")
        return {"records": records, "errors": list(dict.fromkeys(errors)), "warnings": warnings,
                "sheets": sheets, "participants": len(identities), "entries": len(records),
                "duplicates": raw_entries - len(unique)}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "无法读取报名表，请检查文件格式和内容") from exc
    finally:
        if workbook:
            workbook.close()


def validate_import(db: Session, meet_id: int, preview: dict) -> tuple[dict, dict]:
    if not db.get(SportsMeet, meet_id):
        raise HTTPException(404, "运动会不存在")
    events = {event.code: event for event in db.scalars(select(Event).where(Event.meet_id == meet_id, Event.code.in_(ROPE_RULES)))}
    if len(events) != 3:
        raise HTTPException(409, "请先初始化跳绳赛制，再导入报名表")
    teams, grouped = {}, defaultdict(list)
    allowed_colleges = configured_colleges(db)
    for record in preview["records"]:
        if allowed_colleges is not None and record["college"] not in allowed_colleges:
            preview["errors"].append(f"{record['college']}不在学校参赛学院范围内")
        grouped[record["college"]].append(record)
        student = db.scalar(select(StudentRoster).where(StudentRoster.student_no == record["student_no"]))
        if student and (student.name != record["name"] or student.gender != record["gender"] or student.college != record["college"]):
            preview["errors"].append(f"学号{record['student_no']}与已有学生的姓名、性别或学院不一致，请先核查")
    for academy, records in grouped.items():
        matches = list(db.scalars(select(CollegeTeam).where(CollegeTeam.meet_id == meet_id, CollegeTeam.college == academy, CollegeTeam.event_id.is_(None))))
        if len(matches) > 1:
            preview["errors"].append(f"{academy}存在多个通用代表队，请先合并或调整队伍")
            continue
        team = matches[0] if matches else None
        teams[academy] = team
        existing = set(db.scalars(select(StudentRoster.student_no).join(CollegeTeamMember, CollegeTeamMember.student_id == StudentRoster.id).where(CollegeTeamMember.team_id == team.id))) if team else set()
        members = existing | {record["student_no"] for record in records}
        if len(members) > 22:
            preview["errors"].append(f"{academy}导入后总队将有{len(members)}人，超过22人，请核查报名名单")
        elif len(members) < 22:
            preview["warnings"].append(f"{academy}总队当前{len(members)}人，还需补齐至22人")
        for code in {record["event_code"] for record in records if record["event_code"]}:
            lineup = [record for record in records if record["event_code"] == code]
            rule = ROPE_RULES[code]
            jumpers = [record for record in lineup if record["role"] == "JUMPER"]
            turners = [record for record in lineup if record["role"] == "TURNER"]
            if (len(lineup) != rule["lineup_size"] or len(turners) != rule["turner_count"] or
                sum(record["gender"] == "男" for record in jumpers) != rule["male_jumpers"] or
                sum(record["gender"] == "女" for record in jumpers) != rule["female_jumpers"]):
                preview["warnings"].append(f"{academy}的{SHEET_TITLES[code]}阵容需调整为男跳绳者{rule['male_jumpers']}人、女跳绳者{rule['female_jumpers']}人、摇绳者{rule['turner_count']}人")
        codes = {record["event_code"] for record in records if record["event_code"]}
        if team and codes and db.scalar(select(Result.id).join(Registration, Registration.id == Result.registration_id).where(Registration.college_team_id == team.id, Result.event_id.in_([events[code].id for code in codes])).limit(1)):
            preview["errors"].append(f"{academy}已有成绩，不能通过报名导入更换对应项目阵容")
    preview["errors"] = list(dict.fromkeys(preview["errors"]))
    preview["warnings"] = list(dict.fromkeys(preview["warnings"]))
    return events, teams


@router.post("/sports-teacher/rope/import/preview")
def preview_registration(meet_id: int = Form(...), college: str = Form(""), default_event: str = Form(""),
                         last_two_turners: bool = Form(False), file: UploadFile = File(...),
                         _actor: User = Depends(business_admin), db: Session = Depends(get_db)):
    if default_event and default_event not in ROPE_RULES:
        raise HTTPException(422, "默认项目不正确")
    preview = read_registration(file, college, default_event, last_two_turners)
    validate_import(db, meet_id, preview)
    return preview


@router.post("/sports-teacher/rope/import")
def import_registration(meet_id: int = Form(...), college: str = Form(""), default_event: str = Form(""),
                        last_two_turners: bool = Form(False), file: UploadFile = File(...),
                        actor: User = Depends(business_admin), db: Session = Depends(get_db)):
    if default_event and default_event not in ROPE_RULES:
        raise HTTPException(422, "默认项目不正确")
    preview = read_registration(file, college, default_event, last_two_turners)
    events, teams = validate_import(db, meet_id, preview)
    if preview["errors"]:
        raise HTTPException(422, {"message": "报名表校验失败，未导入任何数据", "errors": preview["errors"][:50]})
    created, added = 0, 0
    try:
        students = {}
        lineup_groups = defaultdict(list)
        for record in preview["records"]:
            academy, no = record["college"], record["student_no"]
            team = teams[academy]
            if not team:
                team = CollegeTeam(meet_id=meet_id, college=academy, team_name=f"{academy}代表队", created_by_id=actor.id)
                db.add(team)
                db.flush()
                teams[academy] = team
            student = students.get(no) or db.scalar(select(StudentRoster).where(StudentRoster.student_no == no))
            if not student:
                student = StudentRoster(**{field: record.get(field, "") for field in ("student_no", "name", "gender", "college", "major", "class_name")}, activated=False)
                db.add(student)
                db.flush()
                created += 1
            students[no] = student
            if not db.scalar(select(CollegeTeamMember.id).where(CollegeTeamMember.team_id == team.id, CollegeTeamMember.student_id == student.id)):
                db.add(CollegeTeamMember(team_id=team.id, student_id=student.id, role_title="跳绳队员"))
                db.flush()
                added += 1
            if record["event_code"]:
                lineup_groups[(academy, record["event_code"])].append((record, student))
        for academy, team in teams.items():
            for event in events.values():
                final = db.scalar(select(EventRound).where(EventRound.event_id == event.id, EventRound.stage == "FINAL"))
                if not db.scalar(select(Registration.id).where(Registration.event_id == event.id, Registration.round_id == final.id, Registration.college_team_id == team.id)):
                    db.add(Registration(event_id=event.id, round_id=final.id, college_team_id=team.id, status="REGISTERED"))
        for (academy, code), members in lineup_groups.items():
            team, event = teams[academy], events[code]
            for row in db.scalars(select(RopeLineupMember).where(RopeLineupMember.event_id == event.id, RopeLineupMember.college_team_id == team.id)):
                db.delete(row)
            db.flush()
            # 与成绩录入的男5人、女5人顺序一致，摇绳者在表格末尾。
            members.sort(key=lambda pair: (pair[0]["role"] == "TURNER", pair[1].gender != "男"))
            sequences = defaultdict(int)
            for record, student in members:
                role = record["role"]
                sequences[role] += 1
                db.add(RopeLineupMember(event_id=event.id, college_team_id=team.id, student_id=student.id, role=role, sequence=sequences[role]))
        db.flush()
        for (academy, code) in lineup_groups:
            status = rope_lineup_status(db, events[code], teams[academy].id)
            if not status["valid"]:
                preview["warnings"].append(f"{academy}的{SHEET_TITLES[code]}：" + "；".join(status["errors"]))
        audit(db, actor, "IMPORT_ROPE_REGISTRATION", "sports_meet", meet_id,
              after={"file": file.filename, "participants": preview["participants"], "created": created, "added": added})
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "报名数据已被同时修改，本次导入已取消，请重新预览") from exc
    return {"message": f"报名导入完成：{preview['participants']}人，新增学生{created}人、队员{added}人",
            "created": created, "added": added, "warnings": list(dict.fromkeys(preview["warnings"]))}


def rope_events(db: Session, meet_id: int, actor: User, event_id: int | None = None) -> list[Event]:
    if not db.get(SportsMeet, meet_id):
        raise HTTPException(404, "运动会不存在")
    rows = list(db.scalars(select(Event).where(Event.meet_id == meet_id, Event.code.in_(ROPE_RULES))))
    if event_id:
        rows = [event for event in rows if event.id == event_id]
    if not rows:
        raise HTTPException(404, "该运动会下没有所选跳绳项目")
    rows.sort(key=lambda event: list(ROPE_RULES).index(event.code))
    for event in rows:
        require_event_permission(db, actor, event.id, "can_export")
    return rows


def registrations(db: Session, event: Event) -> list[Registration]:
    return list(db.scalars(select(Registration).join(EventRound, EventRound.id == Registration.round_id)
                          .where(Registration.event_id == event.id, EventRound.stage == "FINAL", Registration.college_team_id.is_not(None))
                          .order_by(Registration.id)))


def ordered_lineup(db: Session, event_id: int, team_id: int):
    rows = [(row, db.get(StudentRoster, row.student_id)) for row in db.scalars(select(RopeLineupMember).where(RopeLineupMember.event_id == event_id, RopeLineupMember.college_team_id == team_id))]
    return sorted(rows, key=lambda pair: (pair[0].role == "TURNER", pair[1].gender != "男", pair[0].sequence, pair[0].id))


def style_sheet(sheet, widths, last_row: int, header_row: int = 4, handwriting: bool = False):
    side = Side(style="thin", color="8A8A8A")
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=header_row, max_row=last_row):
        lines = 1
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.value is not None:
                text_width = sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in str(cell.value))
                lines = max(lines, math.ceil(text_width / max(1, widths[cell.column - 1] - 2)))
        sheet.row_dimensions[row[0].row].height = max(26 if handwriting else 32, lines * 16)
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="E7EDF5")
        cell.font = Font(name="宋体", bold=True, size=11)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":  # 所有输入文本以文本导出，避免公式注入。
                cell.data_type = "s"
    sheet["A1"].font = Font(name="宋体", size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait" if handwriting else "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1 if handwriting else 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.15, footer=.15)
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_area = f"A1:{get_column_letter(len(widths))}{sheet.max_row}"


def download(workbook: Workbook, filename: str, db: Session, actor: User, meet_id: int, action: str):
    stream = BytesIO()
    workbook.save(stream)
    audit(db, actor, action, "sports_meet", meet_id, after={"file": filename})
    db.commit()
    return Response(stream.getvalue(), media_type=MIME, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.get("/admin/rope/record-sheets/export")
def export_record_sheets(meet_id: int, team_id: int | None = None, event_id: int | None = None,
                         actor: User = Depends(manager_only), db: Session = Depends(get_db)):
    events = rope_events(db, meet_id, actor, event_id)
    meet = db.get(SportsMeet, meet_id)
    selected = [(event, registration, db.get(CollegeTeam, registration.college_team_id))
                for event in events for registration in registrations(db, event)
                if not team_id or registration.college_team_id == team_id]
    if not selected:
        raise HTTPException(422, "没有所选学院的跳绳报名名单")
    errors = []
    for event, registration, team in selected:
        status = rope_lineup_status(db, event, team.id)
        if not status["valid"]:
            errors.append(f"{team.college}的{SHEET_TITLES[event.code]}：" + "；".join(status["errors"]))
    if errors:
        raise HTTPException(422, {"message": "请先补齐总队并配置阵容，再导出登记表", "errors": errors[:30]})
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (event, registration, team) in enumerate(selected, 1):
        # 使用短序号保证长学院名称下工作表仍唯一。
        academy_title = re.sub(r"[\\/*?:\[\]]", "_", team.college)[:14]
        sheet = workbook.create_sheet(f"{index:02d}-{SHEET_TITLES[event.code]}-{academy_title}"[:31])
        sheet.merge_cells("A1:F1")
        sheet["A1"] = f"{meet.name}（{meet.edition}）"
        sheet.merge_cells("A2:F2")
        sheet["A2"] = f"跳绳比赛—{SHEET_TITLES[event.code]}成绩登记表"
        sheet["A2"].alignment = Alignment(horizontal="center")
        sheet.merge_cells("A3:F3")
        sheet["A3"] = f"代表队：{team.college} / {team.team_name}    组次：{registration.heat or '待抽签'}    出场序号：{registration.lane or '待抽签'}"
        sheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 30
        sheet.merge_cells("A4:F4")
        sheet["A4"] = f"领队：{team.leader_name or '____________'}    教练：____________"
        sheet.merge_cells("A5:F5")
        sheet["A5"] = "每人1分钟，按男5人、女5人顺序填写次数" if event.code == "ROPE_SINGLE" else "成绩栏填写全队成绩；名单末尾为摇绳者"
        for column, value in enumerate(["序号", "学号", "姓名", "性别", "成绩", "备注"], 1):
            sheet.cell(6, column, value)
        lineup = ordered_lineup(db, event.id, team.id)
        for offset, (member, student) in enumerate(lineup, 7):
            for column, value in enumerate([offset - 6, student.student_no, student.name, student.gender], 1):
                sheet.cell(offset, column, value)
            sheet.cell(offset, 2).number_format = "@"
        end = 6 + len(lineup)
        if event.code != "ROPE_SINGLE":
            sheet.merge_cells(start_row=7, start_column=5, end_row=end, end_column=5)
        sheet.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=3)
        sheet.cell(end + 2, 1, "裁判员（签字）：____________")
        sheet.merge_cells(start_row=end + 2, start_column=4, end_row=end + 2, end_column=6)
        sheet.cell(end + 2, 4, "裁判长（签字）：____________")
        style_sheet(sheet, [8, 24, 15, 9, 18, 24], end, header_row=6, handwriting=True)
    label = selected[0][2].college if team_id else "各学院"
    return download(workbook, f"跳绳_{label}_成绩登记表.xlsx", db, actor, meet_id, "EXPORT_ROPE_RECORD_SHEETS")


def result_text(row: dict | None) -> str:
    if not row:
        return "未发布"
    if row["competition_status"] != "NORMAL":
        return {"DNS": "未参赛", "DNF": "未完成", "DQ": "取消资格"}.get(row["competition_status"], row["competition_status"])
    detail = row.get("rope_detail")
    if detail and row["event_code"] == "ROPE_SYNC":
        return f"完成30次，用时{detail['elapsed_seconds']}秒" if detail["completed_target"] else f"未完成，有效{detail['completed_count']}次"
    if detail:
        return f"{detail['completed_count']}次"
    if row["event_code"] == "ROPE_SYNC" and row["raw_value"] is not None and Decimal(row["raw_value"]) >= 1000:
        return f"未完成，有效{30 - int(Decimal(row['raw_value']) - 1000)}次"
    return f"{row['raw_value']} {row['unit']}"


@router.get("/admin/rope/summary/export")
def export_summary(meet_id: int, actor: User = Depends(manager_only), db: Session = Depends(get_db)):
    events = rope_events(db, meet_id, actor)
    if len(events) != 3:
        raise HTTPException(422, "请先建立完整的三个跳绳项目")
    meet = db.get(SportsMeet, meet_id)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "学院总积分排名"
    totals = {}
    for event in events:
        sheet = workbook.create_sheet(SHEET_TITLES[event.code])
        sheet.merge_cells("A1:G1")
        sheet["A1"] = f"{meet.name} — {SHEET_TITLES[event.code]}最终成绩"
        sheet["A2"] = "使用已发布成绩快照；未发布队伍不计分。并列名次保留。"
        for column, label in enumerate(["名次", "学院", "代表队", "成绩", "积分", "出场序号", "备注"], 1):
            sheet.cell(4, column, label)
        entries = []
        for registration in registrations(db, event):
            team = db.get(CollegeTeam, registration.college_team_id)
            result = db.scalar(select(Result).where(Result.registration_id == registration.id, Result.round_id == registration.round_id, Result.published_at.is_not(None)))
            row = result_dict(db, result, public=True) if result else None
            entries.append((team, registration, row))
            totals.setdefault(team.id, {"team": team, "scores": {}, "published": set()})
            if row:
                totals[team.id]["scores"][event.code] = Decimal(row["points"] or 0)
                totals[team.id]["published"].add(event.code)
        entries.sort(key=lambda item: (item[2] is None, item[2]["rank"] if item[2] and item[2]["rank"] is not None else 999999, item[0].college))
        for team, registration, row in entries:
            sheet.append([row["rank"] if row else None, team.college, team.team_name, result_text(row),
                          float(row["points"]) if row else None, registration.lane, row["notes"] if row else "未发布"])
        style_sheet(sheet, [8, 22, 24, 32, 10, 12, 50], sheet.max_row)
        sheet.auto_filter.ref = f"A4:G{sheet.max_row}"
    summary.merge_cells("A1:H1")
    summary["A1"] = f"{meet.name}（{meet.edition}）跳绳最终汇总"
    summary["A2"] = "按三个小项目已发布积分之和降序排名；未全部发布的学院不列总名次。"
    for column, label in enumerate(["总名次", "学院", "代表队", "计时计数积分", "集体单摇积分", "8字绕绳积分", "总积分", "完成情况"], 1):
        summary.cell(4, column, label)
    ranked = sorted(totals.values(), key=lambda item: (len(item["published"]) != 3, -sum(item["scores"].values(), Decimal(0)), item["team"].college))
    previous, place = None, None
    for index, item in enumerate(ranked, 1):
        complete = len(item["published"]) == 3
        total = sum(item["scores"].values(), Decimal(0))
        if complete and total != previous:
            place, previous = index, total
        summary.append([place if complete else None, item["team"].college, item["team"].team_name,
                        *[float(item["scores"][code]) if code in item["scores"] else None for code in ROPE_RULES],
                        float(total) if item["published"] else None, "全部已发布" if complete else f"已发布{len(item['published'])}/3项"])
    style_sheet(summary, [10, 22, 24, 17, 17, 17, 12, 20], summary.max_row)
    summary.auto_filter.ref = f"A4:H{summary.max_row}"
    return download(workbook, "跳绳_最终成绩排名汇总.xlsx", db, actor, meet_id, "EXPORT_ROPE_SUMMARY")


class RopeDrawRequest(BaseModel):
    meet_id: int
    event_id: int | None = None
    competition_at: datetime
    venue: str = Field(min_length=1, max_length=128)
    slot_minutes: int = Field(default=15, ge=1, le=120)
    check_in_minutes: int = Field(default=15, ge=0, le=120)


@router.post("/sports-teacher/rope/draw")
def draw_schedule(data: RopeDrawRequest, actor: User = Depends(business_admin), db: Session = Depends(get_db)):
    events = rope_events(db, data.meet_id, actor, data.event_id)
    if data.competition_at.tzinfo is not None:
        raise HTTPException(422, "请使用学校当地时间填写比赛时间")
    if not data.venue.strip():
        raise HTTPException(422, "请填写比赛场地")
    if any(db.scalar(select(Result.id).where(Result.event_id == event.id).limit(1)) for event in events):
        raise HTTPException(409, "所选项目已有成绩，不能重新抽签")
    # 对已有赛程不做隐式覆盖，用户可在赛程管理删除草稿后重新生成。
    if any(db.scalar(select(EventSchedule.id).where(EventSchedule.event_id == event.id).limit(1)) for event in events):
        raise HTTPException(409, "所选项目已有赛程，请在赛程管理处理后再抽签")
    groups = []
    errors = []
    for event in events:
        if data.slot_minutes * 60 < ROPE_RULES[event.code]["duration_seconds"]:
            errors.append(f"{SHEET_TITLES[event.code]}每队时段至少{ROPE_RULES[event.code]['duration_seconds'] // 60}分钟")
        rows = [row for row in registrations(db, event) if row.status in {"REGISTERED", "CHECKED_IN"}]
        if not rows:
            errors.append(f"{event.name}没有可抽签的正式报名队伍")
        for row in rows:
            status = rope_lineup_status(db, event, row.college_team_id)
            if not status["valid"]:
                errors.append(f"{db.get(CollegeTeam, row.college_team_id).college}的{SHEET_TITLES[event.code]}阵容未就绪")
        groups.append((event, rows))
    if errors:
        raise HTTPException(422, {"message": "无法生成赛程", "errors": errors})
    current = data.competition_at
    count = 0
    for event, rows in groups:
        SystemRandom().shuffle(rows)
        for order, row in enumerate(rows, 1):
            row.heat, row.lane = f"第{order}组", str(order)
            db.add(EventSchedule(event_id=event.id, round_id=row.round_id, competition_at=current,
                                 check_in_at=current - timedelta(minutes=data.check_in_minutes), venue=data.venue.strip(),
                                 heat=row.heat, notes=f"跳绳抽签：{db.get(CollegeTeam, row.college_team_id).college}",
                                 published=False, status="NOT_STARTED", operator_id=actor.id))
            current += timedelta(minutes=data.slot_minutes)
            count += 1
    audit(db, actor, "DRAW_ROPE_SCHEDULE", "sports_meet", data.meet_id, after={"count": count, "event_ids": [event.id for event in events]})
    db.commit()
    return {"message": f"已随机抽签并生成{count}场赛程草稿，可在赛程管理核对后发布", "count": count}


@router.get("/admin/rope/schedules/export")
def export_schedules(meet_id: int, event_id: int | None = None, actor: User = Depends(manager_only), db: Session = Depends(get_db)):
    events = rope_events(db, meet_id, actor, event_id)
    sheet_book = Workbook()
    sheet = sheet_book.active
    sheet.title = "跳绳抽签赛程"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"{db.get(SportsMeet, meet_id).name}跳绳赛程表"
    sheet["A2"] = "比赛、检录时间按学校当地时间填写；请按发布状态核对。"
    for column, label in enumerate(["项目", "学院", "出场序号", "组次", "比赛时间", "检录时间", "场地", "发布状态"], 1):
        sheet.cell(4, column, label)
    rows = []
    for event in events:
        schedules = list(db.scalars(select(EventSchedule).where(EventSchedule.event_id == event.id)))
        for registration in registrations(db, event):
            for schedule in schedules:
                if schedule.round_id == registration.round_id and schedule.heat == registration.heat:
                    rows.append((event, registration, schedule))
    if not rows:
        raise HTTPException(422, "还没有按抽签生成的跳绳赛程")
    rows.sort(key=lambda item: (item[2].competition_at, item[0].id))
    for event, registration, schedule in rows:
        sheet.append([SHEET_TITLES[event.code], db.get(CollegeTeam, registration.college_team_id).college,
                      int(registration.lane) if (registration.lane or "").isdigit() else registration.lane,
                      registration.heat, schedule.competition_at, schedule.check_in_at, schedule.venue,
                      "已发布" if schedule.published else "草稿"])
        for column in (5, 6):
            sheet.cell(sheet.max_row, column).number_format = "yyyy-mm-dd hh:mm"
    style_sheet(sheet, [22, 22, 12, 12, 24, 24, 24, 12], sheet.max_row)
    return download(sheet_book, "跳绳_抽签赛程表.xlsx", db, actor, meet_id, "EXPORT_ROPE_SCHEDULE")
