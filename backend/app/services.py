from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from fastapi import HTTPException
from fastapi.encoders import jsonable_encoder
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import get_settings
from .models import (
    AdministratorProfile,
    AuditLog,
    CollegeTeam,
    Event,
    EventAdminPermission,
    EventRound,
    Registration,
    Result,
    RopeResultDetail,
    ResultReview,
    ResultState,
    Role,
    ScoreExport,
    ScoringModel,
    ScoringModelBinding,
    ScoringModelItem,
    SportsMeet,
    StudentRoster,
    User,
    UserStatus,
    utcnow,
)


def utc_iso(value: datetime | None) -> str | None:
    """Serialize UTC-naive database timestamps unambiguously for browsers."""
    if value is None:
        return None
    aware = value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
    return aware.isoformat().replace("+00:00", "Z")


ADMIN_ROLES = {Role.SPORTS_TEACHER, Role.REFEREE, Role.SCORER, Role.REFEREE_SCORER}


def snapshot_user(user: User) -> dict:
    return {
        "id": user.id,
        "username": user.username,
        "display_name": user.display_name,
        "role": user.role,
        "department": user.department,
        "status": user.status,
    }


def audit(
    db: Session,
    actor: User,
    action: str,
    target_type: str,
    target_id: int | str | None,
    before: dict | None = None,
    after: dict | None = None,
    reason: str | None = None,
) -> None:
    db.add(
        AuditLog(
            actor_id=actor.id,
            actor_username=actor.username,
            action=action,
            target_type=target_type,
            target_id=str(target_id) if target_id is not None else None,
            before_data=jsonable_encoder(before) if before is not None else None,
            after_data=jsonable_encoder(after) if after is not None else None,
            reason=reason,
        )
    )


def validate_account_quota(db: Session, role: str, exclude_user_id: int | None = None) -> None:
    stmt = select(func.count(User.id))
    if exclude_user_id:
        stmt = stmt.where(User.id != exclude_user_id)
    if role == Role.DEVELOPER:
        count = db.scalar(stmt.where(User.role == Role.DEVELOPER)) or 0
        if count >= 10:
            raise HTTPException(409, "开发人员账号已达到 10 个上限")
    elif role == Role.SPORTS_TEACHER:
        teacher_count = db.scalar(stmt.where(User.role == Role.SPORTS_TEACHER)) or 0
        admin_count = db.scalar(stmt.where(User.role.in_(list(ADMIN_ROLES)))) or 0
        if teacher_count >= 20:
            raise HTTPException(409, "体育部老师账号已达到 20 个上限")
        if admin_count >= 500:
            raise HTTPException(409, "管理人员账号已达到 500 个上限")
    elif role in ADMIN_ROLES:
        admin_count = db.scalar(stmt.where(User.role.in_(list(ADMIN_ROLES)))) or 0
        if admin_count >= 500:
            raise HTTPException(409, "管理人员账号已达到 500 个上限")
    elif role == Role.STUDENT:
        count = db.scalar(stmt.where(User.role == Role.STUDENT)) or 0
        if count >= 10000:
            raise HTTPException(409, "学生账号已达到 10000 个上限")
    else:
        raise HTTPException(422, "无效账号角色")


def ensure_last_developer_safe(db: Session, actor: User, target: User, next_role: str | None = None, freezing: bool = False) -> None:
    removing_active_developer = target.role == Role.DEVELOPER and target.status == UserStatus.ACTIVE and (
        freezing or (next_role is not None and next_role != Role.DEVELOPER)
    )
    if not removing_active_developer:
        return
    active_count = db.scalar(
        select(func.count(User.id)).where(User.role == Role.DEVELOPER, User.status == UserStatus.ACTIVE)
    ) or 0
    if active_count <= 1:
        raise HTTPException(409, "最后一个启用的开发人员账号不能被冻结或降级")
    if actor.id == target.id and active_count <= 1:
        raise HTTPException(409, "不能冻结或降级当前唯一开发人员账号")


def has_event_permission(db: Session, user: User, event_id: int, permission: str) -> bool:
    if user.role in {Role.DEVELOPER, Role.SPORTS_TEACHER}:
        return True
    row = db.scalar(
        select(EventAdminPermission).where(
            EventAdminPermission.user_id == user.id,
            EventAdminPermission.event_id == event_id,
        )
    )
    return bool(row and getattr(row, permission, False))


def require_event_permission(db: Session, user: User, event_id: int, permission: str) -> None:
    if not has_event_permission(db, user, event_id, permission):
        raise HTTPException(403, "没有该比赛项目的操作权限")


def scoring_model_for_event(db: Session, event: Event) -> ScoringModel | None:
    if event.scoring_model_id:
        return db.get(ScoringModel, event.scoring_model_id)
    binding = db.scalar(
        select(ScoringModelBinding)
        .where(ScoringModelBinding.meet_id == event.meet_id, ScoringModelBinding.event_id.is_(None))
        .order_by(ScoringModelBinding.bound_at.desc())
    )
    return db.get(ScoringModel, binding.model_id) if binding else None


def calculate_rankings(db: Session, event_id: int, round_id: int) -> list[Result]:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "比赛项目不存在")
    rows = list(
        db.scalars(
            select(Result).where(Result.event_id == event_id, Result.round_id == round_id)
        )
    )
    normal = [row for row in rows if row.competition_status == "NORMAL" and row.raw_value is not None]
    normal.sort(key=lambda row: row.raw_value, reverse=event.sort_direction == "DESC")
    model = scoring_model_for_event(db, event)
    point_map: dict[int, Decimal] = {}
    if model:
        point_map = {
            item.place: Decimal(item.points)
            for item in db.scalars(select(ScoringModelItem).where(ScoringModelItem.model_id == model.id))
        }
    multiplier = Decimal(event.points_multiplier or 1)
    if model:
        if event.event_type == "RELAY":
            multiplier *= Decimal(model.relay_multiplier)
        elif event.event_type == "TEAM":
            multiplier *= Decimal(model.team_multiplier)
        else:
            multiplier *= Decimal(model.individual_multiplier)

    index = 0
    while index < len(normal):
        tie_end = index + 1
        while tie_end < len(normal) and normal[tie_end].raw_value == normal[index].raw_value:
            tie_end += 1
        place = index + 1
        occupied = list(range(place, tie_end + 1))
        avg = sum((point_map.get(rank, Decimal("0")) for rank in occupied), Decimal("0")) / len(occupied)
        for row in normal[index:tie_end]:
            row.rank = place
            row.points = (avg * multiplier).quantize(Decimal("0.01"))
        index = tie_end

    normal_ids = {row.id for row in normal}
    for row in rows:
        if row.id not in normal_ids:
            row.rank = None
            row.points = Decimal("0")
    db.flush()
    return rows


def add_review(db: Session, result: Result, actor: User, action: str, to_state: str, comment: str | None = None) -> None:
    db.add(
        ResultReview(
            result_id=result.id,
            action=action,
            from_state=result.state,
            to_state=to_state,
            comment=comment,
            operator_id=actor.id,
            operator_name_snapshot=actor.display_name,
            operator_username_snapshot=actor.username,
        )
    )
    result.state = to_state


def public_value(result: Result, field: str):
    published_field = {
        "raw_value": "published_raw_value",
        "rank": "published_rank",
        "points": "published_points",
    }[field]
    if result.published_at is not None:
        return getattr(result, published_field)
    return getattr(result, field)


def mask_student_no(student_no: str) -> str:
    if len(student_no) <= 4:
        return "*" * len(student_no)
    return student_no[:2] + "*" * (len(student_no) - 4) + student_no[-2:]


def sync_event_rounds(db: Session, event_id: int | None = None) -> int:
    event_stmt = select(Event)
    if event_id is not None:
        event_stmt = event_stmt.where(Event.id == event_id)
    changed = 0
    for event in db.scalars(event_stmt):
        existing = list(db.scalars(select(EventRound).where(EventRound.event_id == event.id).order_by(EventRound.id)))
        canonical = {
            "HEAT": (1, "初赛", event.qualification_count),
            "SEMIFINAL": (2, "晋级赛", event.qualification_count),
            "FINAL": (3, "决赛", 0),
        }
        if len(existing) == 3 and all(
            row.stage in canonical
            and (row.sequence, row.name, row.qualification_count) == canonical[row.stage]
            for row in existing
        ):
            continue
        for index, row in enumerate(existing, 1):
            temporary_sequence = 10000 + row.id + index
            if row.sequence != temporary_sequence:
                row.sequence = temporary_sequence
                changed += 1
        if existing:
            db.flush()
        by_stage = {row.stage: row for row in existing}
        for sequence, stage, name in [(1, "HEAT", "初赛"), (2, "SEMIFINAL", "晋级赛"), (3, "FINAL", "决赛")]:
            row = by_stage.get(stage)
            if not row:
                row = EventRound(event_id=event.id, stage=stage, name=name, sequence=sequence, qualification_count=event.qualification_count if stage != "FINAL" else 0)
                db.add(row)
                by_stage[stage] = row
                changed += 1
            else:
                values = {"sequence": sequence, "name": name, "qualification_count": event.qualification_count if stage != "FINAL" else 0}
                for field, value in values.items():
                    if getattr(row, field) != value:
                        setattr(row, field, value)
                        changed += 1
        db.flush()
    return changed


def participant_for_registration(db: Session, registration: Registration | None) -> dict:
    if not registration:
        return {"participant_type": "UNKNOWN", "id": None, "student_no": "", "name": "", "college": "", "class_name": ""}
    if registration.college_team_id is not None:
        team = db.get(CollegeTeam, registration.college_team_id)
        return {
            "participant_type": "TEAM",
            "id": team.id if team else None,
            "student_no": "",
            "name": team.team_name if team else "",
            "college": team.college if team else "",
            "class_name": "学院集体",
        }
    student = db.get(StudentRoster, registration.student_id) if registration.student_id is not None else None
    return {
        "participant_type": "STUDENT",
        "id": student.id if student else None,
        "student_no": student.student_no if student else "",
        "name": student.name if student else "",
        "college": student.college if student else "",
        "class_name": student.class_name if student else "",
        "activated": student.activated if student else False,
    }


def result_dict(db: Session, result: Result, public: bool = False) -> dict:
    registration = db.get(Registration, result.registration_id)
    participant = participant_for_registration(db, registration)
    event = db.get(Event, result.event_id)
    round_row = db.get(EventRound, result.round_id) if result.round_id else None
    from .rope_rules import rope_detail_payload

    rope_detail = db.scalar(select(RopeResultDetail).where(RopeResultDetail.result_id == result.id))
    return {
        "id": result.id,
        "event_id": result.event_id,
        "meet_id": event.meet_id if event else None,
        "round_id": result.round_id,
        "registration_id": result.registration_id,
        "event_name": event.name if event else "",
        "event_code": event.code if event else "",
        "stage": round_row.stage if round_row else (event.stage if event else ""),
        "round_name": round_row.name if round_row else "",
        "unit": event.result_unit if event else "",
        "participant_type": participant["participant_type"],
        "student_id": participant["id"] if participant["participant_type"] == "STUDENT" else None,
        "college_team_id": participant["id"] if participant["participant_type"] == "TEAM" else None,
        "student_no": mask_student_no(participant["student_no"]) if public and participant["student_no"] else participant["student_no"],
        "name": participant["name"],
        "college": participant["college"],
        "class_name": participant["class_name"],
        "heat": registration.heat if registration else None,
        "lane": registration.lane if registration else None,
        "raw_value": public_value(result, "raw_value") if public else result.raw_value,
        "competition_status": result.published_competition_status if public and result.published_at is not None else result.competition_status,
        "rank": public_value(result, "rank") if public else result.rank,
        "points": public_value(result, "points") if public else result.points,
        "state": result.state,
        "published_at": utc_iso(result.published_at),
        "entered_by_id": result.entered_by_id,
        "reviewed_by_id": result.reviewed_by_id,
        "finalized_by_id": result.finalized_by_id,
        "publisher_id": result.publisher_id,
        "entered_at": utc_iso(result.entered_at),
        "reviewed_at": utc_iso(result.reviewed_at),
        "finalized_at": utc_iso(result.finalized_at),
        "version": result.version,
        "published_version": result.published_version,
        "notes": result.published_notes if public and result.published_at is not None else result.notes,
        "rope_detail": rope_detail_payload(rope_detail, public=public, published=result.published_at is not None),
    }


def create_score_workbook(db: Session, event: Event, results: Iterable[Result], exporter: User) -> Path:
    settings = get_settings()
    meet = db.get(SportsMeet, event.meet_id)
    rows = list(results)
    state = rows[0].state if rows else "DRAFT"
    now = utcnow()
    safe_name = f"{meet.name}_{event.name}_{event.stage}_{state}_{now:%Y%m%d_%H%M%S}.xlsx".replace("/", "-")
    export_dir = Path(settings.export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    file_path = export_dir / safe_name

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩单"
    ws.merge_cells("A1:K1")
    ws["A1"] = settings.app_name
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:K2")
    ws["A2"] = f"{meet.name} · {event.name} · {event.stage} · 当前状态：{state}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:K3")
    ws["A3"] = f"导出时间：{now:%Y-%m-%d %H:%M:%S}    导出人：{exporter.display_name}（{exporter.username}）"
    ws["A3"].alignment = Alignment(horizontal="center")
    headers = ["序号", "组次", "道次", "学号/集体编号", "姓名/集体名称", "学院和班级", "原始成绩", "比赛状态", "名次", "积分", "备注"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=value)
        cell.fill = PatternFill("solid", fgColor="DCEBFF")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="9CA3AF")
    for index, result in enumerate(rows, 1):
        registration = db.get(Registration, result.registration_id)
        participant = participant_for_registration(db, registration)
        values = [
            index,
            registration.heat,
            registration.lane,
            mask_student_no(participant["student_no"]) if participant["student_no"] else f"TEAM-{participant['id']}",
            participant["name"],
            f"{participant['college']} / {participant['class_name']}",
            str(result.raw_value or ""),
            result.competition_status,
            result.rank or "",
            float(result.points or 0),
            result.notes or "",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=5 + index, column=col, value=value)

    footer_row = 7 + len(rows)
    action_names = {"ENTER": "成绩登记员", "REVIEW": "成绩审核员", "FINAL_REVIEW": "终审人员", "PUBLISH": "发布人员"}
    for offset, (action, label) in enumerate(action_names.items()):
        review = db.scalar(
            select(ResultReview)
            .join(Result, ResultReview.result_id == Result.id)
            .where(Result.event_id == event.id, ResultReview.action == action)
            .order_by(ResultReview.created_at.desc())
        )
        ws.cell(footer_row + offset, 1, label)
        ws.merge_cells(start_row=footer_row + offset, start_column=2, end_row=footer_row + offset, end_column=11)
        ws.cell(footer_row + offset, 2, f"{review.operator_name_snapshot}（{review.operator_username_snapshot}） {review.created_at:%Y-%m-%d %H:%M:%S}" if review else "")
    ws.cell(footer_row + 3, 1, "本次导出人")
    ws.merge_cells(start_row=footer_row + 3, start_column=2, end_row=footer_row + 3, end_column=11)
    ws.cell(footer_row + 3, 2, f"{exporter.display_name}（{exporter.username}） {now:%Y-%m-%d %H:%M:%S}")

    max_row = footer_row + 3
    for row in ws.iter_rows(min_row=5, max_row=max_row, min_col=1, max_col=11):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [8, 10, 10, 16, 12, 28, 14, 14, 8, 10, 24]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5)
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:K{max_row}"
    ws.freeze_panes = "A6"
    if state != ResultState.PUBLISHED:
        ws.oddHeader.center.text = "&KCCCCCC&24草稿"
    wb.save(file_path)
    db.add(ScoreExport(event_id=event.id, state=state, file_name=safe_name, file_path=str(file_path), exporter_id=exporter.id))
    return file_path
