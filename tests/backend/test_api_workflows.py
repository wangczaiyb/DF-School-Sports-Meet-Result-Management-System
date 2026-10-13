import os
from io import BytesIO
from pathlib import Path


from fastapi.testclient import TestClient
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from backend.app.database import Base, SessionLocal, engine
from backend.app.main import app
from backend.app.models import Role, User
from backend.app.seed import seed_database


@pytest.fixture(scope="session", autouse=True)
def fresh_database():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed_database(db)
    yield
    engine.dispose()


def login(client: TestClient, kind: str, username: str, password: str) -> dict:
    response = client.post(f"/api/v1/auth/{kind}-login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    return response.json()


def headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def test_auth_identity_and_public_visibility():
    with TestClient(app) as client:
        developer = login(client, "developer", "developer01", "Dev123456!")
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        student = login(client, "student", "2026001001", "DF2026001001")
        assert developer["user"]["role"] == "DEVELOPER"
        assert teacher["user"]["role"] == "SPORTS_TEACHER"
        assert student["user"]["role"] == "STUDENT"
        mismatch = client.post("/api/v1/auth/admin-login", json={"username": "developer01", "password": "Dev123456!"})
        assert mismatch.status_code == 403
        public_rows = client.get("/api/v1/public/results").json()
        assert public_rows and all(row["published_at"] for row in public_rows)
        assert all("*" in row["student_no"] for row in public_rows)


def test_account_freeze_unfreeze_and_last_developer_protection():
    with TestClient(app) as client:
        token = login(client, "developer", "developer01", "Dev123456!")["access_token"]
        auth = headers(token)
        blocked = client.post("/api/v1/developer/accounts/1/freeze", headers=auth, json={"reason": "测试"})
        assert blocked.status_code == 409
        created = client.post(
            "/api/v1/developer/accounts",
            headers=auth,
            json={"username": "qa_scorer", "display_name": "测试计分员", "password": "Qa123456!", "role": "SCORER", "department": "体育部", "employee_no": "QA001"},
        )
        assert created.status_code in {200, 409}
        listing = client.get("/api/v1/developer/accounts", headers=auth, params={"keyword": "qa_scorer"}).json()
        target = listing["items"][0]
        assert client.post(f"/api/v1/developer/accounts/{target['id']}/freeze", headers=auth, json={"reason": "流程测试"}).status_code == 200
        assert client.post("/api/v1/auth/admin-login", json={"username": "qa_scorer", "password": "Qa123456!"}).status_code == 403
        assert client.post(f"/api/v1/developer/accounts/{target['id']}/unfreeze", headers=auth).status_code == 200


def test_project_permission_and_complete_result_workflow():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        scorer_referee = login(client, "admin", "admin0004", "Admin123456!")
        scorer = login(client, "admin", "admin0002", "Admin123456!")
        teacher_headers = headers(teacher["access_token"])
        both_headers = headers(scorer_referee["access_token"])
        scorer_headers = headers(scorer["access_token"])
        events = client.get("/api/v1/sports-teacher/events", headers=teacher_headers).json()
        long_jump = next(item for item in events if item["code"] == "WLJ")
        denied = client.get("/api/v1/admin/results", headers=scorer_headers, params={"event_id": long_jump["id"]})
        assert denied.status_code == 403
        registration_data = client.get("/api/v1/admin/registrations", headers=both_headers, params={"event_id": long_jump["id"]}).json()
        registration = registration_data["items"][0]
        round_row = next(row for row in registration_data["rounds"] if row["id"] == registration["round_id"])
        saved = client.put(
            "/api/v1/admin/results",
            headers=both_headers,
            json={"event_id": long_jump["id"], "round_id": round_row["id"], "registration_id": registration["id"], "raw_value": 5.26, "competition_status": "NORMAL"},
        )
        assert saved.status_code == 200, saved.text
        result_id = saved.json()["id"]
        assert client.post("/api/v1/admin/results/submit", headers=both_headers, json={"result_ids": [result_id], "approve": True}).status_code == 200
        assert client.post("/api/v1/admin/results/review", headers=both_headers, json={"result_ids": [result_id], "approve": True, "comment": "成绩有效"}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=teacher_headers, json={"result_ids": [result_id], "approve": True}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/publish", headers=teacher_headers, json={"result_ids": [result_id], "approve": True}).status_code == 200
        public = client.get("/api/v1/public/results", params={"event_id": long_jump["id"]}).json()
        assert any(row["id"] == result_id and float(row["raw_value"]) == 5.26 for row in public)
        exported = client.get("/api/v1/admin/results/export", headers=both_headers, params={"event_id": long_jump["id"]})
        assert exported.status_code == 200
        assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
        workbook = load_workbook(BytesIO(exported.content), read_only=False)
        sheet = workbook["成绩单"]
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.print_area
        assert sheet["A1"].value


def test_account_quota_and_setting_persistence():
    with TestClient(app) as client:
        token = login(client, "developer", "developer01", "Dev123456!")["access_token"]
        auth = headers(token)
        with SessionLocal() as db:
            source = db.scalar(select(User).where(User.username == "developer01"))
            existing = db.scalar(select(User).where(User.username == "quota_dev_2"))
            if not existing:
                for index in range(2, 11):
                    db.add(User(username=f"quota_dev_{index}", password_hash=source.password_hash, role=Role.DEVELOPER, display_name=f"配额测试{index}", status="ACTIVE"))
                db.commit()
        over_limit = client.post("/api/v1/developer/accounts", headers=auth, json={"username": "quota_dev_11", "display_name": "第11个", "password": "Dev123456!", "role": "DEVELOPER"})
        assert over_limit.status_code == 409
        update = client.put("/api/v1/developer/settings", headers=auth, json={"values": {"results_page_name": "测试赛事成绩中心"}})
        assert update.status_code == 200
        assert client.get("/api/v1/public/settings").json()["results_page_name"] == "测试赛事成绩中心"


def test_student_roster_excel_import():
    template = Path("outputs/sports-meet/学生参赛名单导入模板.xlsx")
    assert template.exists()
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        with template.open("rb") as handle:
            response = client.post(
                "/api/v1/sports-teacher/students/import",
                headers=headers(teacher["access_token"]),
                files={"file": (template.name, handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert response.status_code == 200, response.text
        assert response.json()["skipped"] >= 2


def test_manual_student_and_college_team_roster_entry():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        auth = headers(teacher["access_token"])
        student_payload = {
            "student_no": "2026999001",
            "name": "名单测试学生",
            "gender": "女",
            "college": "测试学院",
            "major": "测试专业",
            "class_name": "测试2601",
        }
        created_student = client.post("/api/v1/sports-teacher/students", headers=auth, json=student_payload)
        assert created_student.status_code == 200, created_student.text
        assert created_student.json()["activated"] is False
        duplicate_student = client.post("/api/v1/sports-teacher/students", headers=auth, json=student_payload)
        assert duplicate_student.status_code == 409
        student_list = client.get("/api/v1/sports-teacher/students", headers=auth, params={"keyword": "2026999001"})
        assert student_list.status_code == 200
        assert student_list.json()["total"] == 1
        student_id = student_list.json()["items"][0]["id"]

        meets = client.get("/api/v1/sports-teacher/meets", headers=auth).json()
        meet_id = meets[0]["id"]
        team_payload = {"meet_id": meet_id, "college": "测试学院", "leader_name": "测试领队", "leader_phone": "13800000000"}
        created_team = client.post("/api/v1/sports-teacher/college-teams", headers=auth, json=team_payload)
        assert created_team.status_code == 200, created_team.text
        assert created_team.json()["team_name"] == "测试学院代表队"
        duplicate_team = client.post("/api/v1/sports-teacher/college-teams", headers=auth, json=team_payload)
        assert duplicate_team.status_code == 409
        generated = client.post("/api/v1/sports-teacher/college-teams/generate", headers=auth, json={"meet_id": meet_id})
        assert generated.status_code == 200, generated.text
        teams = client.get("/api/v1/sports-teacher/college-teams", headers=auth, params={"meet_id": meet_id})
        assert teams.status_code == 200
        assert any(item["college"] == "测试学院" for item in teams.json())
        team_id = next(item["id"] for item in teams.json() if item["college"] == "测试学院")

        events = client.get("/api/v1/sports-teacher/events", headers=auth).json()
        female_event = next(item for item in events if item["code"] == "WLJ")
        female_roster = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": female_event["id"]}).json()
        semifinal = next(row for row in female_roster["rounds"] if row["stage"] == "SEMIFINAL")
        final = next(row for row in female_roster["rounds"] if row["stage"] == "FINAL")
        for round_row in [semifinal, final]:
            added = client.post(
                "/api/v1/sports-teacher/event-registrations/batch",
                headers=auth,
                json={"event_id": female_event["id"], "round_id": round_row["id"], "participant_ids": [student_id], "heat": "第1组"},
            )
            assert added.status_code == 200, added.text
            assert added.json()["created"] == 1
        before_activation = client.get("/api/v1/admin/registrations", headers=auth, params={"event_id": female_event["id"]}).json()
        semifinal_registration = next(item for item in before_activation["items"] if item["round_id"] == semifinal["id"] and item["student"]["id"] == student_id)
        inactive_score = client.put(
            "/api/v1/admin/results",
            headers=auth,
            json={"event_id": female_event["id"], "round_id": semifinal["id"], "registration_id": semifinal_registration["id"], "raw_value": 5.18, "competition_status": "NORMAL"},
        )
        assert inactive_score.status_code == 200, inactive_score.text
        activated = client.post("/api/v1/auth/student-register", json={"student_no": student_payload["student_no"], "name": student_payload["name"]})
        assert activated.status_code == 200, activated.text
        student_login = login(client, "student", student_payload["student_no"], "DF2026999001")
        assert student_login["first_login_change_required"] is False
        student_events = client.get("/api/v1/student/my-events", headers=headers(student_login["access_token"])).json()
        matching_rounds = {item["round_name"] for item in student_events if item["event_id"] == female_event["id"]}
        assert {"晋级赛", "决赛"}.issubset(matching_rounds)

        team_event = client.post(
            "/api/v1/sports-teacher/events",
            headers=auth,
            json={
                "meet_id": meet_id,
                "code": "TEAMQA",
                "name": "学院团体测试",
                "gender_group": "MIXED",
                "event_type": "TEAM",
                "result_unit": "分",
                "sort_direction": "DESC",
                "stage": "FINAL",
            },
        )
        assert team_event.status_code == 200, team_event.text
        event_id = team_event.json()["id"]
        event_roster = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": event_id}).json()
        final_round = next(row for row in event_roster["rounds"] if row["stage"] == "FINAL")
        added_team = client.post(
            "/api/v1/sports-teacher/event-registrations/batch",
            headers=auth,
            json={"event_id": event_id, "round_id": final_round["id"], "participant_ids": [team_id], "heat": "决赛"},
        )
        assert added_team.status_code == 200, added_team.text
        team_registrations = client.get("/api/v1/admin/registrations", headers=auth, params={"event_id": event_id})
        assert team_registrations.status_code == 200, team_registrations.text
        registration_data = team_registrations.json()
        assert registration_data["items"]
        assert all(item["student"]["participant_type"] == "TEAM" for item in registration_data["items"])
        saved_team_result = client.put(
            "/api/v1/admin/results",
            headers=auth,
            json={
                "event_id": event_id,
                "round_id": final_round["id"],
                "registration_id": registration_data["items"][0]["id"],
                "raw_value": 88.5,
                "competition_status": "NORMAL",
            },
        )
        assert saved_team_result.status_code == 200, saved_team_result.text
        assert saved_team_result.json()["participant_type"] == "TEAM"
        assert saved_team_result.json()["college"]

        template = client.get("/api/v1/sports-teacher/students/template", headers=auth)
        assert template.status_code == 200
        assert template.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_market_workflow_seeding_advancement_export_schedule_and_crud():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        auth = headers(teacher["access_token"])
        meet_id = client.get("/api/v1/sports-teacher/meets", headers=auth).json()[0]["id"]

        student_ids = []
        for index in range(2):
            payload = {
                "student_no": f"202688880{index}",
                "name": f"晋级测试{index + 1}",
                "gender": "男",
                "college": "流程测试学院",
                "major": "体育测试",
                "class_name": "测试班",
            }
            created = client.post("/api/v1/sports-teacher/students", headers=auth, json=payload)
            assert created.status_code == 200, created.text
            student_ids.append(created.json()["id"])

        event_payload = {
            "meet_id": meet_id,
            "code": "FLOWQA",
            "name": "自动晋级流程测试",
            "gender_group": "MALE",
            "event_type": "INDIVIDUAL",
            "result_unit": "秒",
            "sort_direction": "ASC",
            "stage": "FINAL",
            "qualification_count": 1,
        }
        event_response = client.post("/api/v1/sports-teacher/events", headers=auth, json=event_payload)
        assert event_response.status_code == 200, event_response.text
        event_id = event_response.json()["id"]
        roster = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": event_id}).json()
        heat = next(row for row in roster["rounds"] if row["stage"] == "HEAT")
        semifinal = next(row for row in roster["rounds"] if row["stage"] == "SEMIFINAL")
        added = client.post(
            "/api/v1/sports-teacher/event-registrations/batch",
            headers=auth,
            json={"event_id": event_id, "round_id": heat["id"], "participant_ids": student_ids},
        )
        assert added.status_code == 200 and added.json()["created"] == 2
        seeded = client.post(
            "/api/v1/sports-teacher/event-registrations/seed",
            headers=auth,
            json={"event_id": event_id, "round_id": heat["id"], "lanes_per_heat": 8},
        )
        assert seeded.status_code == 200, seeded.text
        heat_roster = client.get(
            "/api/v1/sports-teacher/event-registrations",
            headers=auth,
            params={"event_id": event_id, "round_id": heat["id"]},
        ).json()["items"]
        assert all(row["heat"] and row["lane"] for row in heat_roster)
        exported = client.get(
            "/api/v1/sports-teacher/event-registrations/export",
            headers=auth,
            params={"event_id": event_id, "round_id": heat["id"]},
        )
        assert exported.status_code == 200
        assert load_workbook(BytesIO(exported.content))["分组表"]["A1"].value == "项目"

        result_ids = []
        for index, registration in enumerate(heat_roster):
            saved = client.put(
                "/api/v1/admin/results",
                headers=auth,
                json={"event_id": event_id, "round_id": heat["id"], "registration_id": registration["id"], "raw_value": 12.1 + index, "competition_status": "NORMAL"},
            )
            assert saved.status_code == 200, saved.text
            result_ids.append(saved.json()["id"])
        assert client.post("/api/v1/admin/results/submit", headers=auth, json={"result_ids": result_ids}).status_code == 200
        assert client.post("/api/v1/admin/results/review", headers=auth, json={"result_ids": result_ids, "approve": True}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=auth, json={"result_ids": result_ids, "approve": True}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/publish", headers=auth, json={"result_ids": result_ids, "approve": True}).status_code == 200
        advanced = client.post(
            "/api/v1/sports-teacher/event-registrations/advance",
            headers=auth,
            json={"event_id": event_id, "from_round_id": heat["id"], "to_round_id": semifinal["id"], "count": 1, "alternates": 1},
        )
        assert advanced.status_code == 200, advanced.text
        semifinal_roster = client.get(
            "/api/v1/sports-teacher/event-registrations",
            headers=auth,
            params={"event_id": event_id, "round_id": semifinal["id"]},
        ).json()["items"]
        assert {row["status"] for row in semifinal_roster} == {"REGISTERED", "ALTERNATE"}

        schedule_payload = {
            "event_id": event_id,
            "round_id": semifinal["id"],
            "competition_at": "2026-10-20T09:00:00",
            "check_in_at": "2026-10-20T08:30:00",
            "venue": "第一田径场",
            "status": "NOT_STARTED",
            "published": True,
        }
        schedule = client.post("/api/v1/admin/schedules", headers=auth, json=schedule_payload)
        assert schedule.status_code == 200, schedule.text
        schedule_payload["venue"] = "主田径场"
        updated = client.put(f"/api/v1/admin/schedules/{schedule.json()['id']}", headers=auth, json=schedule_payload)
        assert updated.status_code == 200 and updated.json()["venue"] == "主田径场"
        assert client.delete(f"/api/v1/admin/schedules/{schedule.json()['id']}", headers=auth).status_code == 200

        model_payload = {
            "name": "流程测试计分",
            "description": "可编辑删除测试",
            "places": 3,
            "points": [5, 3, 1],
            "individual_multiplier": 1,
            "relay_multiplier": 2,
            "team_multiplier": 2,
            "status": "DRAFT",
        }
        model = client.post("/api/v1/sports-teacher/scoring-models", headers=auth, json=model_payload)
        assert model.status_code == 200, model.text
        model_payload["description"] = "已完成修改"
        assert client.put(f"/api/v1/sports-teacher/scoring-models/{model.json()['id']}", headers=auth, json=model_payload).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/scoring-models/{model.json()['id']}", headers=auth).status_code == 200


def test_unpublish_result_workflow():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        auth = headers(teacher["access_token"])
        events = client.get("/api/v1/sports-teacher/events", headers=auth).json()
        target_event = next(item for item in events if item["code"] == "WLJ")
        registration_data = client.get("/api/v1/admin/registrations", headers=auth, params={"event_id": target_event["id"]}).json()
        registration = registration_data["items"][0]
        round_row = next(row for row in registration_data["rounds"] if row["id"] == registration["round_id"])
        
        saved = client.put(
            "/api/v1/admin/results",
            headers=auth,
            json={"event_id": target_event["id"], "round_id": round_row["id"], "registration_id": registration["id"], "raw_value": 4.95, "competition_status": "NORMAL"},
        )
        assert saved.status_code == 200, saved.text
        result_id = saved.json()["id"]
        
        assert client.post("/api/v1/admin/results/submit", headers=auth, json={"result_ids": [result_id]}).status_code == 200
        assert client.post("/api/v1/admin/results/review", headers=auth, json={"result_ids": [result_id], "approve": True}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=auth, json={"result_ids": [result_id], "approve": True}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/publish", headers=auth, json={"result_ids": [result_id], "approve": True}).status_code == 200
        
        pub1 = client.get("/api/v1/public/results", params={"event_id": target_event["id"]}).json()
        assert any(row["id"] == result_id for row in pub1)
        
        unpub = client.post("/api/v1/sports-teacher/results/unpublish", headers=auth, json={"result_ids": [result_id]})
        assert unpub.status_code == 200, unpub.text
        
        pub2 = client.get("/api/v1/public/results", params={"event_id": target_event["id"]}).json()
        assert not any(row["id"] == result_id for row in pub2)


def test_college_teams_per_event_and_member_editing():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        auth = headers(teacher["access_token"])
        meets = client.get("/api/v1/sports-teacher/meets", headers=auth).json()
        meet_id = meets[0]["id"]

        # 创建一个专项比赛项目（如男女混合排球赛）
        ev_resp = client.post(
            "/api/v1/sports-teacher/events",
            headers=auth,
            json={
                "meet_id": meet_id,
                "code": "VB_TEST_TEAM",
                "name": "排球专项测试赛",
                "gender_group": "MIXED",
                "event_type": "TEAM",
                "result_unit": "分",
                "sort_direction": "DESC",
                "stage": "FINAL",
                "qualification_count": 0,
            },
        )
        assert ev_resp.status_code == 200, ev_resp.text
        event_id = ev_resp.json()["id"]

        # 1. 为同一学院（如计算机学院）创建不同比赛项目的代表队
        t1_resp = client.post(
            "/api/v1/sports-teacher/college-teams",
            headers=auth,
            json={
                "meet_id": meet_id,
                "event_id": event_id,
                "college": "计算机学院",
                "team_name": "计算机学院排球测试队",
                "leader_name": "王教练",
                "leader_phone": "13800138000",
            },
        )
        assert t1_resp.status_code == 200, t1_resp.text
        team1 = t1_resp.json()
        assert team1["event_id"] == event_id
        assert team1["event_name"] == "排球专项测试赛"

        # 2. 创建通用代表队
        t2_resp = client.post(
            "/api/v1/sports-teacher/college-teams",
            headers=auth,
            json={
                "meet_id": meet_id,
                "event_id": None,
                "college": "计算机学院",
                "team_name": "计算机学院田径通用代表队",
            },
        )
        assert t2_resp.status_code == 200, t2_resp.text
        team2 = t2_resp.json()
        assert team2["event_id"] is None

        # 3. 筛选代表队测试
        filtered_teams = client.get("/api/v1/sports-teacher/college-teams", headers=auth, params={"meet_id": meet_id, "event_id": event_id}).json()
        assert any(t["id"] == team1["id"] for t in filtered_teams)

        # 4. 创建学生并添加为比赛成员
        s1 = client.post("/api/v1/sports-teacher/students", headers=auth, json={
            "student_no": "2026999901", "name": "李队员", "gender": "女", "college": "计算机学院", "major": "软件工程", "class_name": "2601班"
        }).json()
        s2 = client.post("/api/v1/sports-teacher/students", headers=auth, json={
            "student_no": "2026999902", "name": "张队长", "gender": "男", "college": "计算机学院", "major": "软件工程", "class_name": "2601班"
        }).json()

        # 批量添加队员
        batch_add = client.post(
            f"/api/v1/sports-teacher/college-teams/{team1['id']}/members/batch",
            headers=auth,
            json={"student_ids": [s1["id"], s2["id"]], "role_title": "队员"},
        )
        assert batch_add.status_code == 200, batch_add.text
        assert batch_add.json()["added"] == 2

        # 查询队员列表及达标状态
        members_info = client.get(f"/api/v1/sports-teacher/college-teams/{team1['id']}/members", headers=auth).json()
        assert members_info["total_members"] == 2
        assert members_info["female_count"] == 1
        assert members_info["male_count"] == 1

        # 5. 编辑队员角色与首发状态
        captain_member = next(m for m in members_info["members"] if m["student_id"] == s2["id"])
        update_m = client.put(
            f"/api/v1/sports-teacher/college-teams/{team1['id']}/members/{captain_member['id']}",
            headers=auth,
            json={"role_title": "队长", "is_starter": True},
        )
        assert update_m.status_code == 200, update_m.text
        assert update_m.json()["role_title"] == "队长"
        assert update_m.json()["is_starter"] is True

        # 6. 一键为专项比赛项目批量生成各学院代表队
        gen_resp = client.post(
            "/api/v1/sports-teacher/college-teams/generate",
            headers=auth,
            json={"meet_id": meet_id, "event_id": event_id},
        )
        assert gen_resp.status_code == 200, gen_resp.text


def test_duplicate_meet_and_event_creation_returns_conflict():
    with TestClient(app) as client:
        teacher = login(client, "admin", "admin0001", "Admin123456!")
        auth = headers(teacher["access_token"])

        meet_payload = {
            "name": "  唯一约束回归测试运动会  ",
            "edition": "  2099  ",
            "start_date": "2099-05-01",
            "end_date": "2099-05-02",
        }
        created_meet = client.post("/api/v1/sports-teacher/meets", headers=auth, json=meet_payload)
        assert created_meet.status_code == 200, created_meet.text
        assert created_meet.json()["name"] == "唯一约束回归测试运动会"
        assert created_meet.json()["edition"] == "2099"

        duplicate_meet = client.post(
            "/api/v1/sports-teacher/meets",
            headers=auth,
            json={**meet_payload, "name": "唯一约束回归测试运动会", "edition": "2099"},
        )
        assert duplicate_meet.status_code == 409, duplicate_meet.text
        assert duplicate_meet.json()["detail"] == "同名同届运动会已存在"

        same_edition_meet = client.post(
            "/api/v1/sports-teacher/meets",
            headers=auth,
            json={**meet_payload, "name": "同年度另一类型运动会", "edition": "2099"},
        )
        assert same_edition_meet.status_code == 200, same_edition_meet.text

        event_payload = {
            "meet_id": created_meet.json()["id"],
            "code": "  DUPLICATE_QA  ",
            "name": "  重复编号回归测试项目  ",
            "gender_group": "MALE",
        }
        created_event = client.post("/api/v1/sports-teacher/events", headers=auth, json=event_payload)
        assert created_event.status_code == 200, created_event.text
        assert created_event.json()["code"] == "DUPLICATE_QA"
        assert created_event.json()["name"] == "重复编号回归测试项目"

        duplicate_event = client.post(
            "/api/v1/sports-teacher/events",
            headers=auth,
            json={**event_payload, "code": "DUPLICATE_QA", "name": "另一个项目名称"},
        )
        assert duplicate_event.status_code == 409, duplicate_event.text
        assert duplicate_event.json()["detail"] == "该运动会中项目编号已存在"

        same_code_other_meet = client.post(
            "/api/v1/sports-teacher/events",
            headers=auth,
            json={**event_payload, "meet_id": same_edition_meet.json()["id"], "code": "DUPLICATE_QA"},
        )
        assert same_code_other_meet.status_code == 200, same_code_other_meet.text

        blank_event = client.post(
            "/api/v1/sports-teacher/events",
            headers=auth,
            json={**event_payload, "code": "   "},
        )
        assert blank_event.status_code == 422, blank_event.text
