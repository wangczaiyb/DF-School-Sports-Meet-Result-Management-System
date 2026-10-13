import os
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO


import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.exc import IntegrityError

from backend.app.database import Base, SessionLocal, engine
from backend.app.main import app
from backend.app.models import (
    AdministratorProfile,
    AuditLog,
    CollegeTeam,
    CollegeTeamMember,
    Event,
    EventRound,
    EventSchedule,
    Registration,
    Result,
    ResultReview,
    ScoringModel,
    SportsMeet,
    StudentRoster,
    SystemSetting,
    User,
    utcnow,
)
from backend.app.security import verify_password
from backend.app.seed import seed_database


@pytest.fixture(scope="module", autouse=True)
def fresh_verification_database():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed_database(db)
    yield


def _login(client: TestClient, kind: str, username: str, password: str) -> dict:
    response = client.post(f"/api/v1/auth/{kind}-login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    return response.json()


def _headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _teacher_headers(client: TestClient) -> dict:
    return _headers(_login(client, "admin", "admin0001", "Admin123456!")["access_token"])


def _event_payload(meet_id: int, code: str, name: str, event_type: str = "INDIVIDUAL", **overrides) -> dict:
    payload = {
        "meet_id": meet_id,
        "code": code,
        "name": name,
        "gender_group": "MIXED" if event_type != "INDIVIDUAL" else "MALE",
        "participant_group": "STUDENT",
        "event_type": event_type,
        "result_unit": "次" if event_type != "INDIVIDUAL" else "秒",
        "sort_direction": "DESC" if event_type != "INDIVIDUAL" else "ASC",
        "stage": "FINAL",
        "qualification_count": 0,
        "scoring_places": 8,
        "points_multiplier": 1,
    }
    payload.update(overrides)
    return payload


def _student_payload(index: int, *, college: str = "CODEX_TEST_学院", gender: str = "男") -> dict:
    return {
        "student_no": f"CODEX_TEST_{index:06d}",
        "name": f"CODEX_TEST_学生_{index}",
        "gender": gender,
        "college": college,
        "major": "CODEX_TEST_专业",
        "class_name": "CODEX_TEST_班级",
    }


def test_auth_account_profile_lock_force_logout_and_setting_persistence():
    with TestClient(app) as client:
        developer = _login(client, "developer", "developer01", "Dev123456!")
        auth = _headers(developer["access_token"])

        created = client.post(
            "/api/v1/developer/accounts",
            headers=auth,
            json={
                "username": "CODEX_TEST_scorer_account",
                "display_name": "CODEX_TEST_计分员",
                "password": "CodexTest123!",
                "role": "SCORER",
                "department": "CODEX_TEST_体育部",
                "employee_no": "CODEX_TEST_EMP_001",
                "phone": "13800000001",
            },
        )
        assert created.status_code == 200, created.text
        account_id = created.json()["id"]
        assert created.json()["employee_no"] == "CODEX_TEST_EMP_001"

        with SessionLocal() as db:
            user = db.get(User, account_id)
            profile = db.scalar(select(AdministratorProfile).where(AdministratorProfile.user_id == account_id))
            assert user and user.password_hash != "CodexTest123!"
            assert verify_password("CodexTest123!", user.password_hash)
            assert profile and profile.employee_no == "CODEX_TEST_EMP_001" and profile.phone == "13800000001"

        updated = client.put(
            f"/api/v1/developer/accounts/{account_id}",
            headers=auth,
            json={
                "username": "CODEX_TEST_scorer_account_v2",
                "display_name": "CODEX_TEST_计分员_已更新",
                "department": "CODEX_TEST_竞赛部",
                "employee_no": "CODEX_TEST_EMP_002",
                "phone": "13800000002",
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["employee_no"] == "CODEX_TEST_EMP_002"
        listing = client.get("/api/v1/developer/accounts", headers=auth, params={"keyword": "CODEX_TEST_scorer_account_v2"}).json()
        assert listing["total"] == 1
        assert listing["items"][0]["phone"] == "13800000002"

        scorer_login = _login(client, "admin", "CODEX_TEST_scorer_account_v2", "CodexTest123!")
        scorer_auth = _headers(scorer_login["access_token"])
        assert client.post(f"/api/v1/developer/accounts/{account_id}/force-logout", headers=auth).status_code == 200
        assert client.get("/api/v1/auth/me", headers=scorer_auth).status_code == 401

        assert client.post(f"/api/v1/developer/accounts/{account_id}/freeze", headers=auth, json={"reason": "CODEX_TEST_冻结"}).status_code == 200
        assert client.post("/api/v1/auth/admin-login", json={"username": "CODEX_TEST_scorer_account_v2", "password": "CodexTest123!"}).status_code == 403
        assert client.post(f"/api/v1/developer/accounts/{account_id}/unfreeze", headers=auth).status_code == 200
        assert client.post(f"/api/v1/developer/accounts/{account_id}/reset-password", headers=auth, json={"password": "CodexReset123!"}).status_code == 200
        _login(client, "admin", "CODEX_TEST_scorer_account_v2", "CodexReset123!")

        locked = client.post(
            "/api/v1/developer/accounts",
            headers=auth,
            json={
                "username": "CODEX_TEST_lock_account",
                "display_name": "CODEX_TEST_锁定账号",
                "password": "CodexLock123!",
                "role": "SCORER",
                "employee_no": "CODEX_TEST_LOCK_001",
            },
        )
        assert locked.status_code == 200, locked.text
        for _ in range(5):
            wrong = client.post("/api/v1/auth/admin-login", json={"username": "CODEX_TEST_lock_account", "password": "wrong-password"})
            assert wrong.status_code == 401
        assert client.post("/api/v1/auth/admin-login", json={"username": "CODEX_TEST_lock_account", "password": "CodexLock123!"}).status_code == 423
        with SessionLocal() as db:
            locked_user = db.scalar(select(User).where(User.username == "CODEX_TEST_lock_account"))
            assert locked_user.locked_until and locked_user.locked_until >= utcnow() + timedelta(minutes=14)
            locked_user.locked_until = utcnow() - timedelta(seconds=1)
            db.commit()
        _login(client, "admin", "CODEX_TEST_lock_account", "CodexLock123!")

        settings_update = client.put(
            "/api/v1/developer/settings",
            headers=auth,
            json={"values": {"results_page_name": "CODEX_TEST_赛事成绩中心"}},
        )
        assert settings_update.status_code == 200
        assert client.get("/api/v1/public/settings").json()["results_page_name"] == "CODEX_TEST_赛事成绩中心"
        with SessionLocal() as db:
            assert db.scalar(select(SystemSetting.value).where(SystemSetting.key == "results_page_name")) == "CODEX_TEST_赛事成绩中心"
            actions = set(db.scalars(select(AuditLog.action).where(AuditLog.target_id == str(account_id))))
            assert {"CREATE_ACCOUNT", "UPDATE_ACCOUNT", "FREEZE_ACCOUNT", "UNFREEZE_ACCOUNT", "RESET_PASSWORD", "FORCE_LOGOUT"}.issubset(actions)


def test_meet_event_student_team_schedule_crud_and_cross_meet_isolation():
    with TestClient(app) as client:
        auth = _teacher_headers(client)
        meet_payload = {
            "name": "CODEX_TEST_春季运动会",
            "edition": "2035",
            "start_date": "2035-04-01",
            "end_date": "2035-04-03",
            "status": "PREPARING",
            "public_visible": True,
        }
        meet_a = client.post("/api/v1/sports-teacher/meets", headers=auth, json=meet_payload)
        meet_b = client.post("/api/v1/sports-teacher/meets", headers=auth, json={**meet_payload, "name": "CODEX_TEST_趣味运动会"})
        assert meet_a.status_code == meet_b.status_code == 200
        meet_a_id, meet_b_id = meet_a.json()["id"], meet_b.json()["id"]
        meet_payload["name"] = "CODEX_TEST_春季田径运动会_已更新"
        assert client.put(f"/api/v1/sports-teacher/meets/{meet_a_id}", headers=auth, json=meet_payload).status_code == 200

        event_a = client.post("/api/v1/sports-teacher/events", headers=auth, json=_event_payload(meet_a_id, "CODEX_TEST_SHARED", "CODEX_TEST_个人赛A"))
        event_b = client.post("/api/v1/sports-teacher/events", headers=auth, json=_event_payload(meet_b_id, "CODEX_TEST_SHARED", "CODEX_TEST_个人赛B"))
        team_event_1 = client.post("/api/v1/sports-teacher/events", headers=auth, json=_event_payload(meet_a_id, "CODEX_TEST_TEAM_A", "CODEX_TEST_团体赛A", "TEAM"))
        team_event_2 = client.post("/api/v1/sports-teacher/events", headers=auth, json=_event_payload(meet_a_id, "CODEX_TEST_TEAM_B", "CODEX_TEST_团体赛B", "TEAM"))
        assert all(response.status_code == 200 for response in [event_a, event_b, team_event_1, team_event_2])
        event_a_id, event_b_id = event_a.json()["id"], event_b.json()["id"]
        team_event_1_id, team_event_2_id = team_event_1.json()["id"], team_event_2.json()["id"]
        assert {row["id"] for row in client.get("/api/v1/sports-teacher/events", headers=auth, params={"meet_id": meet_a_id}).json()} == {event_a_id, team_event_1_id, team_event_2_id}
        assert {row["id"] for row in client.get("/api/v1/sports-teacher/events", headers=auth, params={"meet_id": meet_b_id}).json()} == {event_b_id}

        cross_team = client.post(
            "/api/v1/sports-teacher/college-teams",
            headers=auth,
            json={"meet_id": meet_a_id, "event_id": event_b_id, "college": "CODEX_TEST_学院", "team_name": "CODEX_TEST_跨届队"},
        )
        assert cross_team.status_code == 422

        student = client.post("/api/v1/sports-teacher/students", headers=auth, json=_student_payload(100001))
        assert student.status_code == 200, student.text
        student_id = student.json()["id"]
        updated_student_payload = {**_student_payload(100001), "name": "CODEX_TEST_学生_已更新", "major": "CODEX_TEST_新专业"}
        assert client.put(f"/api/v1/sports-teacher/students/{student_id}", headers=auth, json=updated_student_payload).status_code == 200
        assert client.post("/api/v1/sports-teacher/students", headers=auth, json=updated_student_payload).status_code == 409
        assert client.get("/api/v1/sports-teacher/students", headers=auth, params={"keyword": "CODEX_TEST_100001"}).json()["total"] == 1

        team = client.post(
            "/api/v1/sports-teacher/college-teams",
            headers=auth,
            json={"meet_id": meet_a_id, "event_id": team_event_1_id, "college": "CODEX_TEST_学院", "team_name": "CODEX_TEST_专项队"},
        )
        assert team.status_code == 200, team.text
        team_id = team.json()["id"]
        member = client.post(
            f"/api/v1/sports-teacher/college-teams/{team_id}/members",
            headers=auth,
            json={"student_id": student_id, "role_title": "CODEX_TEST_队长", "is_starter": True},
        )
        assert member.status_code == 200, member.text
        assert client.delete(f"/api/v1/sports-teacher/students/{student_id}", headers=auth).status_code == 409
        assert client.delete(f"/api/v1/sports-teacher/college-teams/{team_id}", headers=auth).status_code == 409

        roster2 = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": team_event_2_id}).json()
        assert team_id not in {row["id"] for row in roster2["candidates"]}
        team2_final = next(row for row in roster2["rounds"] if row["stage"] == "FINAL")
        wrong_project = client.post(
            "/api/v1/sports-teacher/event-registrations/batch",
            headers=auth,
            json={"event_id": team_event_2_id, "round_id": team2_final["id"], "participant_ids": [team_id]},
        )
        assert wrong_project.status_code == 422
        cross_update = client.put(
            f"/api/v1/sports-teacher/college-teams/{team_id}",
            headers=auth,
            json={"meet_id": meet_a_id, "event_id": event_b_id, "college": "CODEX_TEST_学院", "team_name": "CODEX_TEST_专项队"},
        )
        assert cross_update.status_code == 422

        roster_a = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": event_a_id}).json()
        heat = next(row for row in roster_a["rounds"] if row["stage"] == "HEAT")
        added = client.post(
            "/api/v1/sports-teacher/event-registrations/batch",
            headers=auth,
            json={"event_id": event_a_id, "round_id": heat["id"], "participant_ids": [student_id]},
        )
        assert added.status_code == 200
        registration = client.get("/api/v1/sports-teacher/event-registrations", headers=auth, params={"event_id": event_a_id, "round_id": heat["id"]}).json()["items"][0]
        assert client.put(f"/api/v1/sports-teacher/event-registrations/{registration['id']}", headers=auth, json={"status": "INVALID_STATUS"}).status_code == 422
        changed = client.put(
            f"/api/v1/sports-teacher/event-registrations/{registration['id']}",
            headers=auth,
            json={"bib_no": "CODEX_TEST_BIB", "heat": "第1组", "lane": "4", "status": "CHECKED_IN"},
        )
        assert changed.status_code == 200 and changed.json()["status"] == "CHECKED_IN"
        with SessionLocal() as db:
            persisted = db.get(Registration, registration["id"])
            assert persisted.status == "CHECKED_IN" and persisted.heat == "第1组" and persisted.lane == "4"

        invalid_schedule = client.post(
            "/api/v1/admin/schedules",
            headers=auth,
            json={"event_id": event_a_id, "round_id": heat["id"], "competition_at": "2035-04-01T09:00:00", "check_in_at": "2035-04-01T09:30:00", "venue": "CODEX_TEST_场地", "published": True},
        )
        assert invalid_schedule.status_code == 422
        assert client.post(
            "/api/v1/admin/schedules",
            headers=auth,
            json={"event_id": 999999999, "competition_at": "2035-04-01T09:00:00", "venue": "CODEX_TEST_场地"},
        ).status_code == 404
        schedule_payload = {"event_id": event_a_id, "round_id": heat["id"], "competition_at": "2035-04-01T09:00:00", "check_in_at": "2035-04-01T08:30:00", "venue": "CODEX_TEST_场地A", "published": True}
        schedule = client.post("/api/v1/admin/schedules", headers=auth, json=schedule_payload)
        assert schedule.status_code == 200, schedule.text
        schedule_id = schedule.json()["id"]
        schedule_payload["venue"] = "CODEX_TEST_场地B"
        assert client.put(f"/api/v1/admin/schedules/{schedule_id}", headers=auth, json=schedule_payload).json()["venue"] == "CODEX_TEST_场地B"
        assert any(row["id"] == schedule_id and row["venue"] == "CODEX_TEST_场地B" for row in client.get("/api/v1/public/schedules").json())
        assert client.delete(f"/api/v1/admin/schedules/{schedule_id}", headers=auth).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/event-registrations/{registration['id']}", headers=auth).status_code == 200

        assert client.delete(f"/api/v1/sports-teacher/college-teams/{team_id}/members/{member.json()['id']}", headers=auth).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/students/{student_id}", headers=auth).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/college-teams/{team_id}", headers=auth).status_code == 200
        for event_id in [event_a_id, team_event_1_id, team_event_2_id, event_b_id]:
            assert client.delete(f"/api/v1/sports-teacher/events/{event_id}", headers=auth).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/meets/{meet_a_id}", headers=auth).status_code == 200
        assert client.delete(f"/api/v1/sports-teacher/meets/{meet_b_id}", headers=auth).status_code == 200
        with SessionLocal() as db:
            assert db.scalar(select(func.count(SportsMeet.id)).where(SportsMeet.name.like("CODEX_TEST_%"))) == 0


def test_result_state_machine_ties_snapshot_promotion_audit_and_permissions():
    from backend.app.models import SportsMeet

    with TestClient(app) as client:
        teacher_auth = _teacher_headers(client)
        meet = client.post(
            "/api/v1/sports-teacher/meets",
            headers=teacher_auth,
            json={"name": "CODEX_TEST_成绩闭环运动会", "edition": "2036", "start_date": "2036-05-01", "end_date": "2036-05-03"},
        )
        assert meet.status_code == 200
        meet_id = meet.json()["id"]
        model = client.post(
            "/api/v1/sports-teacher/scoring-models",
            headers=teacher_auth,
            json={"name": "CODEX_TEST_并列计分模型", "places": 8, "points": [9, 7, 6, 5, 4, 3, 2, 1], "individual_multiplier": 1, "relay_multiplier": 2, "team_multiplier": 2, "status": "ENABLED"},
        )
        assert model.status_code == 200, model.text
        event = client.post(
            "/api/v1/sports-teacher/events",
            headers=teacher_auth,
            json=_event_payload(meet_id, "CODEX_TEST_RESULT", "CODEX_TEST_成绩闭环项目", qualification_count=2),
        )
        assert event.status_code == 200, event.text
        event_id = event.json()["id"]
        assert client.post(f"/api/v1/sports-teacher/scoring-models/{model.json()['id']}/bind", headers=teacher_auth, json={"meet_id": meet_id, "event_id": event_id}).status_code == 200

        student_ids = []
        for index in range(3):
            created = client.post("/api/v1/sports-teacher/students", headers=teacher_auth, json=_student_payload(200001 + index, college="CODEX_TEST_成绩学院"))
            assert created.status_code == 200, created.text
            student_ids.append(created.json()["id"])
        roster = client.get("/api/v1/sports-teacher/event-registrations", headers=teacher_auth, params={"event_id": event_id}).json()
        heat = next(row for row in roster["rounds"] if row["stage"] == "HEAT")
        semifinal = next(row for row in roster["rounds"] if row["stage"] == "SEMIFINAL")
        assert client.post("/api/v1/sports-teacher/event-registrations/batch", headers=teacher_auth, json={"event_id": event_id, "round_id": heat["id"], "participant_ids": student_ids}).status_code == 200
        assert client.post("/api/v1/sports-teacher/event-registrations/seed", headers=teacher_auth, json={"event_id": event_id, "round_id": heat["id"], "lanes_per_heat": 8}).status_code == 200
        heat_items = client.get("/api/v1/sports-teacher/event-registrations", headers=teacher_auth, params={"event_id": event_id, "round_id": heat["id"]}).json()["items"]
        assert all(item["heat"] and item["lane"] for item in heat_items)

        result_ids = []
        for registration, raw in zip(heat_items, [10.0, 10.0, 11.0]):
            saved = client.put(
                "/api/v1/admin/results",
                headers=teacher_auth,
                json={"event_id": event_id, "round_id": heat["id"], "registration_id": registration["id"], "raw_value": raw, "competition_status": "NORMAL", "notes": "CODEX_TEST_原始公开备注"},
            )
            assert saved.status_code == 200, saved.text
            result_ids.append(saved.json()["id"])
        ranked = client.get("/api/v1/admin/results", headers=teacher_auth, params={"event_id": event_id, "round_id": heat["id"]}).json()
        assert [row["rank"] for row in ranked] == [1, 1, 3]
        assert [float(row["points"]) for row in ranked] == [8.0, 8.0, 6.0]

        assert client.post("/api/v1/sports-teacher/results/final-review", headers=teacher_auth, json={"result_ids": [result_ids[0]], "approve": True}).status_code == 409
        assert client.post("/api/v1/sports-teacher/results/publish", headers=teacher_auth, json={"result_ids": [result_ids[0]], "approve": True}).status_code == 409
        assert client.post("/api/v1/admin/results/submit", headers=teacher_auth, json={"result_ids": [result_ids[0]]}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=teacher_auth, json={"result_ids": [result_ids[0]], "approve": True}).status_code == 409
        rejected = client.post("/api/v1/admin/results/review", headers=teacher_auth, json={"result_ids": [result_ids[0]], "approve": False, "comment": "CODEX_TEST_退回"})
        assert rejected.status_code == 200
        assert next(row for row in client.get("/api/v1/admin/results", headers=teacher_auth, params={"event_id": event_id}).json() if row["id"] == result_ids[0])["state"] == "RETURNED"
        edited = client.put(
            "/api/v1/admin/results",
            headers=teacher_auth,
            json={"event_id": event_id, "round_id": heat["id"], "registration_id": heat_items[0]["id"], "raw_value": 10.0, "competition_status": "NORMAL", "notes": "CODEX_TEST_原始公开备注"},
        )
        assert edited.status_code == 200
        assert client.post("/api/v1/admin/results/submit", headers=teacher_auth, json={"result_ids": result_ids}).status_code == 200
        assert client.post("/api/v1/admin/results/review", headers=teacher_auth, json={"result_ids": result_ids, "approve": True, "comment": "CODEX_TEST_裁判通过"}).status_code == 200
        assert client.post("/api/v1/sports-teacher/results/publish", headers=teacher_auth, json={"result_ids": result_ids, "approve": True}).status_code == 409
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=teacher_auth, json={"result_ids": result_ids, "approve": True, "comment": "CODEX_TEST_终审通过"}).status_code == 200
        with SessionLocal() as db:
            finalized = [db.get(Result, result_id) for result_id in result_ids]
            assert all(row.finalized_by_id and row.finalized_at and row.publisher_id is None for row in finalized)
        assert client.post("/api/v1/sports-teacher/results/publish", headers=teacher_auth, json={"result_ids": result_ids, "approve": True, "comment": "CODEX_TEST_发布"}).status_code == 200

        public_before = client.get("/api/v1/public/results", params={"event_id": event_id}).json()
        original_public = next(row for row in public_before if row["id"] == result_ids[0])
        assert float(original_public["raw_value"]) == 10.0 and original_public["competition_status"] == "NORMAL"
        assert original_public["published_at"].endswith("Z")
        changed_internal = client.put(
            "/api/v1/admin/results",
            headers=teacher_auth,
            json={"event_id": event_id, "round_id": heat["id"], "registration_id": heat_items[0]["id"], "raw_value": 9.0, "competition_status": "DNF", "notes": "CODEX_TEST_未发布修改", "modification_reason": "CODEX_TEST_快照验证"},
        )
        assert changed_internal.status_code == 200
        assert changed_internal.json()["state"] == "DRAFT" and float(changed_internal.json()["raw_value"]) == 9.0
        public_after = client.get("/api/v1/public/results", params={"event_id": event_id}).json()
        unchanged_public = next(row for row in public_after if row["id"] == result_ids[0])
        assert float(unchanged_public["raw_value"]) == 10.0
        assert unchanged_public["competition_status"] == "NORMAL"
        assert unchanged_public["notes"] == "CODEX_TEST_原始公开备注"
        public_meets = client.get("/api/v1/public/meets").json()
        assert any(row["id"] == meet_id for row in public_meets)
        scoped_public = client.get("/api/v1/public/results", params={"meet_id": meet_id}).json()
        assert scoped_public and all(row["meet_id"] == meet_id for row in scoped_public)
        other_meet_id = next(row["id"] for row in public_meets if row["id"] != meet_id)
        assert not any(row["id"] in result_ids for row in client.get("/api/v1/public/results", params={"meet_id": other_meet_id}).json())
        assert any(row["college"] == "CODEX_TEST_成绩学院" for row in client.get("/api/v1/public/college-rankings", params={"meet_id": meet_id}).json())
        assert not any(row["college"] == "CODEX_TEST_成绩学院" for row in client.get("/api/v1/public/college-rankings", params={"meet_id": other_meet_id}).json())

        advanced = client.post(
            "/api/v1/sports-teacher/event-registrations/advance",
            headers=teacher_auth,
            json={"event_id": event_id, "from_round_id": heat["id"], "to_round_id": semifinal["id"], "count": 2, "alternates": 1},
        )
        assert advanced.status_code == 200, advanced.text
        next_round = client.get("/api/v1/sports-teacher/event-registrations", headers=teacher_auth, params={"event_id": event_id, "round_id": semifinal["id"]}).json()["items"]
        assert [row["status"] for row in next_round].count("REGISTERED") == 2
        assert [row["status"] for row in next_round].count("ALTERNATE") == 1

        activated = client.post("/api/v1/auth/student-register", json={"student_no": "CODEX_TEST_200001", "name": "CODEX_TEST_学生_200001"})
        assert activated.status_code == 200, activated.text
        student_login = _login(client, "student", "CODEX_TEST_200001", "DFCODEX_TEST_200001")
        student_auth = _headers(student_login["access_token"])
        assert any(row["event_id"] == event_id for row in client.get("/api/v1/student/my-results", headers=student_auth).json())
        assert any(row["event_id"] == event_id and row["stage"] == "SEMIFINAL" for row in client.get("/api/v1/student/my-qualifications", headers=student_auth).json())

        scorer_auth = _headers(_login(client, "admin", "admin0002", "Admin123456!")["access_token"])
        referee_auth = _headers(_login(client, "admin", "admin0003", "Admin123456!")["access_token"])
        both_auth = _headers(_login(client, "admin", "admin0004", "Admin123456!")["access_token"])
        assert client.get("/api/v1/admin/results", headers=scorer_auth, params={"event_id": event_id}).status_code == 403
        assert client.post("/api/v1/admin/results/submit", headers=scorer_auth, json={"result_ids": [result_ids[0]]}).status_code == 403
        assert client.post("/api/v1/sports-teacher/results/final-review", headers=scorer_auth, json={"result_ids": [result_ids[0]], "approve": True}).status_code == 403
        assert client.get("/api/v1/developer/accounts", headers=both_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/meets", headers=student_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/basketball/teams").status_code == 401
        assert client.get("/api/v1/sports-teacher/basketball/teams", headers=scorer_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/basketball/teams", headers=teacher_auth).status_code == 200
        assert client.get("/api/v1/sports-teacher/volleyball/teams").status_code == 401
        assert client.get("/api/v1/sports-teacher/volleyball/teams", headers=scorer_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/volleyball/teams", headers=teacher_auth).status_code == 200
        assert client.get("/api/v1/sports-teacher/table-tennis/team/teams").status_code == 401
        assert client.get("/api/v1/sports-teacher/table-tennis/team/teams", headers=scorer_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/table-tennis/team/teams", headers=teacher_auth, params={"event_id": event_id}).status_code == 200
        assert client.get("/api/v1/sports-teacher/table-tennis/singles/athletes").status_code == 401
        assert client.get("/api/v1/sports-teacher/table-tennis/singles/athletes", headers=scorer_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/table-tennis/singles/athletes", headers=teacher_auth, params={"event_id": event_id}).status_code == 200
        assert client.get("/api/v1/sports-teacher/rope/status").status_code == 401
        assert client.get("/api/v1/sports-teacher/rope/status", headers=scorer_auth).status_code == 403
        assert client.get("/api/v1/sports-teacher/rope/status", headers=teacher_auth, params={"meet_id": meet_id}).status_code == 200

        seed_events = client.get("/api/v1/sports-teacher/events", headers=teacher_auth).json()
        m100 = next(row for row in seed_events if row["code"] == "M100")
        wlj = next(row for row in seed_events if row["code"] == "WLJ")
        assert client.get("/api/v1/admin/results", headers=scorer_auth, params={"event_id": m100["id"]}).status_code == 200
        assert client.get("/api/v1/admin/results", headers=scorer_auth, params={"event_id": wlj["id"]}).status_code == 403
        assert client.get("/api/v1/admin/results/export", headers=scorer_auth, params={"event_id": m100["id"]}).status_code == 200
        assert client.get("/api/v1/admin/results/export", headers=scorer_auth, params={"event_id": wlj["id"]}).status_code == 403
        assert client.post("/api/v1/admin/results/review", headers=referee_auth, json={"result_ids": [result_ids[0]], "approve": True}).status_code == 403

        with SessionLocal() as db:
            published = db.get(Result, result_ids[1])
            assert published.publisher_id and published.finalized_by_id and published.publisher_id == published.finalized_by_id
            assert published.published_competition_status == "NORMAL"
            review_actions = set(db.scalars(select(ResultReview.action).where(ResultReview.result_id == result_ids[0])))
            assert {"ENTER", "RETURN", "REVIEW", "FINAL_REVIEW", "PUBLISH"}.issubset(review_actions)
            audit_actions = set(db.scalars(select(AuditLog.action).where(AuditLog.target_type == "result", AuditLog.target_id.in_([str(x) for x in result_ids]))))
            assert {"SAVE_RESULT", "SUBMIT_RESULT", "RETURN_RESULT", "REVIEW_RESULT", "FINAL_REVIEW_RESULT", "PUBLISH_RESULT"}.issubset(audit_actions)
            scoring = db.get(ScoringModel, model.json()["id"])
            assert scoring.locked_by_published_results is True


def test_excel_atomic_import_exports_sqlite_pragmas_foreign_keys_and_concurrency():
    def workbook_bytes(rows: list[list[str]]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "学生名单"
        sheet.append(["学号", "姓名", "性别", "学院", "专业", "班级"])
        for row in rows:
            sheet.append(row)
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    with TestClient(app) as client:
        auth = _teacher_headers(client)
        atomic_rows = [
            ["CODEX_TEST_300001", "CODEX_TEST_原子学生1", "男", "CODEX_TEST_导入学院", "CODEX_TEST_专业", "CODEX_TEST_班"],
            ["CODEX_TEST_300002", "CODEX_TEST_原子学生2", "未知", "CODEX_TEST_导入学院", "CODEX_TEST_专业", "CODEX_TEST_班"],
        ]
        invalid = client.post(
            "/api/v1/sports-teacher/students/import",
            headers=auth,
            files={"file": ("CODEX_TEST_atomic.xlsx", workbook_bytes(atomic_rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert invalid.status_code == 422
        with SessionLocal() as db:
            assert db.scalar(select(func.count(StudentRoster.id)).where(StudentRoster.student_no.in_(["CODEX_TEST_300001", "CODEX_TEST_300002"]))) == 0

        valid_rows = [
            ["CODEX_TEST_300011", "CODEX_TEST_导入学生1", "男", "CODEX_TEST_导入学院", "CODEX_TEST_专业", "CODEX_TEST_班"],
            ["CODEX_TEST_300012", "CODEX_TEST_导入学生2", "女", "CODEX_TEST_导入学院", "CODEX_TEST_专业", "CODEX_TEST_班"],
        ]
        valid_bytes = workbook_bytes(valid_rows)
        imported = client.post(
            "/api/v1/sports-teacher/students/import",
            headers=auth,
            files={"file": ("CODEX_TEST_valid.xlsx", valid_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 200 and imported.json()["created"] == 2
        duplicate = client.post(
            "/api/v1/sports-teacher/students/import",
            headers=auth,
            files={"file": ("CODEX_TEST_duplicate.xlsx", valid_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert duplicate.status_code == 200 and duplicate.json()["created"] == 0 and duplicate.json()["skipped"] == 2
        assert client.post("/api/v1/sports-teacher/students/import", headers=auth, files={"file": ("CODEX_TEST_wrong.csv", b"a,b", "text/csv")}).status_code == 422
        assert client.post("/api/v1/sports-teacher/students/import", headers=auth, files={"file": ("CODEX_TEST_empty.xlsx", b"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code == 422
        template = client.get("/api/v1/sports-teacher/students/template", headers=auth)
        assert template.status_code == 200 and template.headers["content-type"].startswith("application/vnd.openxmlformats")

        scorer_auth = _headers(_login(client, "admin", "admin0002", "Admin123456!")["access_token"])
        event = next(row for row in client.get("/api/v1/admin/my-events", headers=scorer_auth).json() if row["event"]["code"] == "M100")
        export = client.get("/api/v1/admin/results/export", headers=scorer_auth, params={"event_id": event["event"]["id"]})
        assert export.status_code == 200
        workbook = load_workbook(BytesIO(export.content))
        sheet = workbook["成绩单"]
        assert sheet.page_setup.orientation == "landscape"
        assert str(sheet.page_setup.paperSize) == str(sheet.PAPERSIZE_A4)
        assert [sheet.cell(5, column).value for column in range(1, 12)] == ["序号", "组次", "道次", "学号/集体编号", "姓名/集体名称", "学院和班级", "原始成绩", "比赛状态", "名次", "积分", "备注"]

    with engine.connect() as connection:
        assert connection.execute(text("PRAGMA foreign_keys")).scalar() == 1
        assert connection.execute(text("PRAGMA journal_mode")).scalar().lower() == "wal"
        assert connection.execute(text("PRAGMA busy_timeout")).scalar() == 30000
        assert connection.execute(text("PRAGMA integrity_check")).scalar() == "ok"
        assert connection.execute(text("PRAGMA foreign_key_check")).fetchall() == []
        connection.rollback()
        transaction = connection.begin()
        try:
            with pytest.raises(IntegrityError):
                connection.execute(text("INSERT INTO college_team_members(team_id, student_id, is_starter) VALUES (999999999, 999999999, 0)"))
        finally:
            transaction.rollback()

    def concurrent_write(index: int) -> str:
        key = f"CODEX_TEST_concurrent_{index}"
        with SessionLocal() as db:
            db.add(SystemSetting(key=key, value=f"value-{index}"))
            db.commit()
        return key

    with ThreadPoolExecutor(max_workers=3) as executor:
        keys = list(executor.map(concurrent_write, range(3)))
    with SessionLocal() as db:
        values = dict(db.execute(select(SystemSetting.key, SystemSetting.value).where(SystemSetting.key.in_(keys))).all())
        assert values == {f"CODEX_TEST_concurrent_{index}": f"value-{index}" for index in range(3)}
        assert db.scalar(select(func.count(AuditLog.id)).where(AuditLog.action == "IMPORT_STUDENT_ROSTER")) >= 2
