"""报名识别、学院登记表、抽签赛程和发布快照汇总的端到端验证。"""
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from backend.app.database import SessionLocal
from backend.app.main import app
from backend.app.models import (CollegeTeam, CollegeTeamMember, EventSchedule, Result,
                                EventAdminPermission, RopeLineupMember, StudentRoster, User)


@pytest.fixture
def rope_context():
    client = TestClient(app)
    login = client.post('/api/v1/auth/admin-login', json={'username': 'admin0001', 'password': 'Admin123456!'})
    headers = {'Authorization': f"Bearer {login.json()['access_token']}"}
    meet = client.post('/api/v1/sports-teacher/meets', headers=headers, json={
        'name': f'表格验证{uuid4().hex[:12]}', 'edition': '2026',
        'start_date': '2026-10-20', 'end_date': '2026-10-21', 'status': 'PREPARING',
    })
    assert meet.status_code == 200, meet.text
    meet_id = meet.json()['id']
    initialized = client.post('/api/v1/sports-teacher/rope/initialize', headers=headers, json={'meet_id': meet_id})
    assert initialized.status_code == 200, initialized.text
    return client, headers, meet_id


def workbook_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def registration_bytes(prefix, college, explicit_roles=True):
    book = Workbook()
    book.remove(book.active)
    male = [(f'{prefix}1{i:02d}', f'男队员{i}', '男') for i in range(1, 12)]
    female = [(f'{prefix}2{i:02d}', f'女队员{i}', '女') for i in range(1, 12)]
    for title, jumpers, turners in [
        ('计时计数项目', male[:5] + female[:5], []),
        ('集体单摇跳', male[:5] + female[:5], male[-1:] + female[-1:]),
        ('“8”字绕绳跳', male[:10] + female[:10], male[-1:] + female[-1:]),
    ]:
        sheet = book.create_sheet(title)
        sheet.append(['运动会报名表'])
        sheet.append([f'代表队：{college}'])
        sheet.append(['序号', '学生学号', '运动员姓名', '男女', '角色', '成绩', '备注'])
        for index, (no, name, gender) in enumerate(jumpers + turners, 1):
            sheet.append([index, no, name, gender, ('跳绳者' if index <= len(jumpers) else '摇绳者') if explicit_roles else None, 999, '忽略报名表成绩'])
    return workbook_bytes(book)


def upload(ctx, content, path='import', **fields):
    client, headers, meet_id = ctx
    return client.post(f'/api/v1/sports-teacher/rope/{path}', headers=headers,
                       data={'meet_id': str(meet_id), **fields}, files={'file': ('报名.xlsx', content)})


def status(ctx):
    client, headers, meet_id = ctx
    return client.get('/api/v1/sports-teacher/rope/status', headers=headers, params={'meet_id': meet_id}).json()


def test_download_template_can_import_roster_and_both_role_sections(rope_context):
    template = Path(__file__).resolve().parents[2] / 'frontend/src/assets/rope-registration-template.xlsx'
    book = load_workbook(template)
    assert book.sheetnames == ['运动员汇总表（22人）', '1.前单跳', '2.集体单摇跳', '3.“8”字环绕跳']
    male = list(range(1, 12))
    female = list(range(12, 23))
    turners = [11, 22]
    groups = [male + female, male[:5] + female[:5], male[:5] + female[:5] + turners,
              male[:10] + female[:10] + turners]
    prefix = f'8{rope_context[2]:04d}'
    for sheet, members in zip(book, groups):
        rows = [row for row in range(3, sheet.max_row + 1) if isinstance(sheet.cell(row, 1).value, int)]
        assert len(rows) == len(members)
        for row, person in zip(rows, members):
            assert all(sheet.cell(row, col).value is None for col in range(2, 6))
            assert sheet.cell(row, 3).number_format == '@'
            sheet.cell(row, 2, '模板验证学院')
            sheet.cell(row, 3, f'{prefix}{person:03d}')
            sheet.cell(row, 4, f'队员{person}')
            sheet.cell(row, 5, '男' if person in male else '女')
    content = workbook_bytes(book)
    response = upload(rope_context, content, path='import/preview')
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['participants'] == 22 and preview['entries'] == 66
    assert not preview['errors'] and not preview['warnings']
    for code in ('ROPE_SYNC', 'ROPE_FIG8'):
        assert {row['student_no'] for row in preview['records'] if row['event_code'] == code and row['role'] == 'TURNER'} == {
            f'{prefix}{person:03d}' for person in turners}
    response = upload(rope_context, content)
    assert response.status_code == 200, response.text
    assert not response.json()['warnings']
    assert all(event['teams'][0]['lineup_status']['valid'] for event in status(rope_context)['events'])


def test_import_preview_and_college_print_sheets(rope_context):
    ctx = rope_context
    content = registration_bytes(f'1{ctx[2]:04d}', '表格甲学院')
    with SessionLocal() as db:
        before = db.scalar(select(func.count(StudentRoster.id)))
    response = upload(ctx, content, path='import/preview')
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['participants'] == 22 and preview['entries'] == 44
    assert not preview['errors'] and not preview['warnings']
    with SessionLocal() as db:
        assert db.scalar(select(func.count(StudentRoster.id))) == before
    for _ in range(2):
        response = upload(ctx, content)
        assert response.status_code == 200, response.text
        assert not response.json()['warnings']
    imported = status(ctx)
    team_id = imported['events'][0]['teams'][0]['team_id']
    with SessionLocal() as db:
        assert db.scalar(select(func.count(CollegeTeamMember.id)).where(CollegeTeamMember.team_id == team_id)) == 22
        assert db.scalar(select(func.count(RopeLineupMember.id)).where(RopeLineupMember.college_team_id == team_id)) == 44
        assert db.scalar(select(func.count(Result.id)).where(Result.event_id.in_([event['id'] for event in imported['events']]))) == 0
    assert all(event['teams'][0]['lineup_status']['valid'] for event in imported['events'])
    other = upload(ctx, registration_bytes(f'2{ctx[2]:04d}', '表格乙学院'))
    assert other.status_code == 200, other.text
    client, headers, meet_id = ctx
    exported = client.get('/api/v1/admin/rope/record-sheets/export', headers=headers, params={'meet_id': meet_id, 'team_id': team_id})
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content))
    assert len(book.worksheets) == 3
    for sheet, size in zip(book, [10, 12, 22]):
        assert sheet['A3'].value.startswith('代表队：表格甲学院')
        assert '表格乙学院' not in str(list(sheet.values))
        assert [sheet.cell(6, column).value for column in range(1, 7)] == ['序号', '学号', '姓名', '性别', '成绩', '备注']
        assert sheet['B7'].value.startswith(f'1{meet_id:04d}')
        assert all(sheet.cell(row, 2).data_type == 's' for row in range(7, 7 + size))
        assert all(sheet.cell(row, column).value is None for row in range(7, 7 + size) for column in (5, 6))
        assert all(sheet.cell(row, column).value for row in range(7, 7 + size) for column in (2, 3, 4))
        assert sheet.page_setup.fitToHeight == 1
        assert len(sheet.print_area) > 0


def test_missing_roles_need_explicit_choice_and_conflicts_are_atomic(rope_context):
    ctx = rope_context
    content = registration_bytes(f'3{ctx[2]:04d}', '角色学院', explicit_roles=False)
    preview = upload(ctx, content, path='import/preview').json()
    assert not preview['errors'] and preview['warnings']
    explicit = upload(ctx, content, path='import/preview', last_two_turners='true').json()
    assert not explicit['errors'] and not explicit['warnings']
    assert upload(ctx, content, last_two_turners='true').status_code == 200
    book = load_workbook(BytesIO(content))
    book.worksheets[0]['C4'] = '同学号冲突姓名'
    with SessionLocal() as db:
        before = db.scalar(select(func.count(StudentRoster.id)))
    response = upload(ctx, workbook_bytes(book), last_two_turners='true')
    assert response.status_code == 422
    with SessionLocal() as db:
        assert db.scalar(select(func.count(StudentRoster.id))) == before
    empty = Workbook()
    empty.active.append(['序号', '学号', '姓名', '性别', '成绩', '备注'])
    assert upload(ctx, workbook_bytes(empty), path='import/preview').json()['errors']


def test_merged_college_multi_project_csv_and_text_ids(rope_context):
    ctx = rope_context
    book = Workbook()
    sheet = book.active
    sheet.append(['学院', '学号', '姓名', '性别', '参赛项目'])
    sheet.append(['合并学院', '00123456', '甲同学', '男性', '前单跳、集体单摇跳'])
    sheet.append([None, '00123457', '乙同学', 'F', '8字绕绳跳'])
    sheet.merge_cells('A2:A3')
    preview = upload(ctx, workbook_bytes(book), path='import/preview').json()
    assert not preview['errors']
    assert preview['participants'] == 2 and preview['entries'] == 3
    assert all(record['college'] == '合并学院' for record in preview['records'])
    assert preview['records'][0]['student_no'] == '00123456'
    source = '学号,姓名,性别,前单跳,集体单摇跳,8字绕绳跳\n00123458,丙同学,女,√,,是\n'.encode('gb18030')
    client, headers, meet_id = ctx
    response = client.post('/api/v1/sports-teacher/rope/import/preview', headers=headers,
                           data={'meet_id': meet_id, 'college': 'CSV学院'}, files={'file': ('报名.csv', source)})
    assert response.status_code == 200, response.text
    assert len(response.json()['records']) == 2
    assert not response.json()['errors']


@pytest.mark.parametrize('number', ['1234567', '123456789', 'ABCDEFGH', '1234_678', '１２３４５６７８'])
def test_rope_import_rejects_non_eight_digit_numbers_atomically(rope_context, number):
    book = Workbook()
    book.active.append(['学号', '姓名', '性别', '学院'])
    book.active.append(['00123459', '有效学生', '男', '校验学院'])
    book.active.append([number, '无效学生', '女', '校验学院'])
    content = workbook_bytes(book)
    preview = upload(rope_context, content, path='import/preview').json()
    assert any('8位数字' in error for error in preview['errors'])
    with SessionLocal() as db:
        before = db.scalar(select(func.count(StudentRoster.id)))
    assert upload(rope_context, content).status_code == 422
    with SessionLocal() as db:
        assert db.scalar(select(func.count(StudentRoster.id))) == before


def test_draw_generates_nonoverlapping_drafts_and_exports(rope_context):
    ctx = rope_context
    for suffix, college in [('A', '抽签甲学院'), ('B', '抽签乙学院')]:
        assert upload(ctx, registration_bytes(f'4{ctx[2]:03d}{ord(suffix) - ord("A")}', college)).status_code == 200
    client, headers, meet_id = ctx
    payload = {'meet_id': meet_id, 'competition_at': '2026-10-20T09:00:00', 'venue': '体育馆', 'slot_minutes': 15}
    response = client.post('/api/v1/sports-teacher/rope/draw', headers=headers, json=payload)
    assert response.status_code == 200, response.text
    assert response.json()['count'] == 6
    with SessionLocal() as db:
        ids = [event['id'] for event in status(ctx)['events']]
        rows = list(db.scalars(select(EventSchedule).where(EventSchedule.event_id.in_(ids)).order_by(EventSchedule.competition_at)))
        assert len(rows) == 6 and all(not row.published for row in rows)
        assert all((second.competition_at - first.competition_at).total_seconds() == 900 for first, second in zip(rows, rows[1:]))
        assert all((row.competition_at - row.check_in_at).total_seconds() == 900 for row in rows)
    assert client.post('/api/v1/sports-teacher/rope/draw', headers=headers, json=payload).status_code == 409
    exported = client.get('/api/v1/admin/rope/schedules/export', headers=headers, params={'meet_id': meet_id})
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    assert book.active.max_row == 10
    assert all(book.active.cell(row, 8).value == '草稿' for row in range(5, 11))


def test_summary_uses_published_snapshots_ties_and_incomplete_status(rope_context):
    ctx = rope_context
    for suffix, college in [('A', '排名甲学院'), ('B', '排名乙学院'), ('C', '未发布学院')]:
        assert upload(ctx, registration_bytes(f'5{ctx[2]:03d}{ord(suffix) - ord("A")}', college)).status_code == 200
    client, headers, meet_id = ctx
    result_ids = []
    for event in status(ctx)['events']:
        for team in event['teams'][:2]:
            payload = {'event_id': event['id'], 'round_id': event['round_id'], 'registration_id': team['registration_id']}
            payload.update({'individual_counts': [100] * 10} if event['code'] == 'ROPE_SINGLE' else
                           {'completed_target': True, 'elapsed_seconds': 45} if event['code'] == 'ROPE_SYNC' else {'completed_count': 100})
            response = client.put('/api/v1/admin/rope/results', headers=headers, json=payload)
            assert response.status_code == 200, response.text
            result_ids.append(response.json()['id'])
    body = {'result_ids': result_ids, 'approve': True}
    for path in ['admin/results/submit', 'admin/results/review', 'sports-teacher/results/final-review', 'sports-teacher/results/publish']:
        response = client.post(f'/api/v1/{path}', headers=headers, json=body)
        assert response.status_code == 200, response.text
    # 草稿数据发生变化，汇总仍须保留发布快照。
    with SessionLocal() as db:
        result = db.get(Result, result_ids[0])
        result.raw_value = 9999
        result.points = 99
        result.rank = 9
        db.commit()
    response = client.get('/api/v1/admin/rope/summary/export', headers=headers, params={'meet_id': meet_id})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    summary = book.worksheets[0]
    assert summary['A5'].value == summary['A6'].value == 1
    assert summary['G5'].value == summary['G6'].value == 24
    assert summary['A7'].value is None and summary['H7'].value == '已发布0/3项'
    assert book.worksheets[1]['D5'].value == '1000次'
    assert book.worksheets[1]['E5'].value == 8
    # 无项目导出权限的计分员不得获取其他项目人员信息。
    login = client.post('/api/v1/auth/admin-login', json={'username': 'admin0002', 'password': 'Admin123456!'})
    restricted = {'Authorization': f"Bearer {login.json()['access_token']}"}
    assert client.get('/api/v1/admin/rope/summary/export', headers=restricted, params={'meet_id': meet_id}).status_code == 403
    assert client.get('/api/v1/admin/rope/status', headers=restricted, params={'meet_id': meet_id}).status_code == 403
    event = status(ctx)['events'][0]
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == 'admin0002'))
        teacher = db.scalar(select(User).where(User.username == 'admin0001'))
        db.add(EventAdminPermission(user_id=user.id, event_id=event['id'], can_export=True, assigned_by_id=teacher.id))
        db.commit()
    permitted = client.get('/api/v1/admin/rope/status', headers=restricted, params={'meet_id': meet_id})
    assert permitted.status_code == 200
    assert [row['id'] for row in permitted.json()['events']] == [event['id']]
    meets = client.get('/api/v1/admin/rope/meets', headers=restricted).json()
    assert any(row['id'] == meet_id for row in meets)
    assert client.get('/api/v1/admin/rope/record-sheets/export', headers=restricted,
                      params={'meet_id': meet_id, 'event_id': event['id']}).status_code == 200
