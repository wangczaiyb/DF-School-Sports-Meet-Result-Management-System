"""验证花名册批量注册、事务回滚及统一学生初始密码。"""

from io import BytesIO

from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from backend.app.api import (
    _login, create_account, download_student_account_template,
    import_student_accounts, reset_password, router, student_register,
)
from backend.app.database import Base
from backend.app.dependencies import developer_only, get_db
from backend.app.models import AuditLog, Role, StudentRoster, User
from backend.app.schemas import AccountCreate, LoginRequest, ResetPasswordRequest, StudentRegisterRequest
from backend.app.security import hash_password, verify_password


@pytest.fixture(scope="module")
def password_hash():
    return hash_password("Dev123456!")


@pytest.fixture
def database(tmp_path, password_hash):
    engine = create_engine(f"sqlite:///{(tmp_path / 'students.db').as_posix()}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        actor = User(username="developer01", display_name="开发人员", role=Role.DEVELOPER,
                     password_hash=password_hash, status="ACTIVE")
        db.add(actor)
        db.commit()
        yield db, actor
    engine.dispose()


def upload(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生名单"
    sheet.append(headers or ["学号", "姓名"])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return UploadFile(file=stream, filename="花名册.xlsx")


def student(db, student_no):
    return db.scalar(select(User).where(User.username == student_no))


def test_two_columns_register_accounts_and_allow_login(database):
    db, actor = database
    result = import_student_accounts(upload([["00123456", "学生甲"], [2026999002, "学生乙"]]), actor, db)
    assert result["created"] == result["registered"] == 2
    assert result["skipped"] == 0
    for student_no in ("00123456", "2026999002"):
        account = student(db, student_no)
        assert account.role == Role.STUDENT
        assert verify_password(f"DF{student_no}", account.password_hash)
        assert not verify_password(f"DF{student_no[-6:]}", account.password_hash)
        roster = db.scalar(select(StudentRoster).where(StudentRoster.student_no == student_no))
        assert roster.activated and roster.user_id == account.id
        login = _login(db, LoginRequest(username=student_no, password=f"DF{student_no}"), {Role.STUDENT})
        assert login["first_login_change_required"] is False
    assert db.scalar(select(func.count(AuditLog.id)).where(AuditLog.action == "CREATE_ACCOUNT")) == 2


def test_optional_columns_and_existing_unactivated_roster(database):
    db, actor = database
    roster = StudentRoster(student_no="2026999010", name="待注册学生", gender="女",
                           college="现有学院", major="现有专业", class_name="现有班级", activated=False)
    db.add(roster)
    db.commit()
    original_id = roster.id
    result = import_student_accounts(upload([
        ["待注册学生", "2026999010", "女", "导入学院"],
        ["新学生", "2026999011", "男", "导入学院"],
    ], ["姓名*", "学号*", "性别", "学院"]), actor, db)
    assert result["registered"] == 2 and result["created"] == 1
    db.refresh(roster)
    assert roster.id == original_id and roster.activated and roster.college == "现有学院"
    added = db.scalar(select(StudentRoster).where(StudentRoster.student_no == "2026999011"))
    assert added.gender == "男" and added.college == "导入学院"


def test_reimport_preserves_existing_password(database):
    db, actor = database
    rows = [["2026999020", "重复学生"], ["2026999020", "重复学生"]]
    result = import_student_accounts(upload(rows), actor, db)
    assert result["registered"] == 1 and result["skipped"] == 1
    account = student(db, "2026999020")
    account.password_hash = hash_password("Changed123456!")
    db.commit()
    original_hash = account.password_hash
    result = import_student_accounts(upload(rows), actor, db)
    assert result["created"] == result["registered"] == 0 and result["skipped"] == 2
    db.refresh(account)
    assert account.password_hash == original_hash


@pytest.mark.parametrize("rows", [
    [["2026999030", "有效学生"], ["2026999031", ""]],
    [["2026999030", "有效学生"], ["2026999030", "同学号不同姓名"]],
])
def test_invalid_batch_is_atomic(database, rows):
    db, actor = database
    with pytest.raises(HTTPException) as error:
        import_student_accounts(upload(rows), actor, db)
    assert error.value.status_code == 422
    assert db.scalar(select(func.count(StudentRoster.id))) == 0
    assert db.scalar(select(func.count(User.id)).where(User.role == Role.STUDENT)) == 0


def test_account_name_collision_is_atomic(database):
    db, actor = database
    actor.username = "2026999040"
    db.commit()
    with pytest.raises(HTTPException) as error:
        import_student_accounts(upload([["2026999041", "有效学生"], ["2026999040", "冲突学生"]]), actor, db)
    assert error.value.status_code == 422
    assert student(db, "2026999041") is None


def test_quota_blocks_whole_batch(database, password_hash):
    db, actor = database
    db.execute(User.__table__.insert(), [
        {"username": f"quota_{index:06d}", "display_name": "配额学生", "role": Role.STUDENT,
         "password_hash": password_hash, "status": "ACTIVE"}
        for index in range(10000)
    ])
    db.commit()
    with pytest.raises(HTTPException) as error:
        import_student_accounts(upload([["2026999050", "超过配额学生"]]), actor, db)
    assert error.value.status_code == 409
    assert student(db, "2026999050") is None
    assert db.scalar(select(func.count(StudentRoster.id))) == 0


def test_developer_permission_and_template(database):
    db, actor = database
    teacher = User(username="teacher", display_name="老师", role=Role.SPORTS_TEACHER,
                   password_hash=actor.password_hash, status="ACTIVE")
    with pytest.raises(HTTPException) as error:
        developer_only(teacher)
    assert error.value.status_code == 403
    for path in ("/developer/students/import", "/developer/students/template"):
        route = next(route for route in router.routes if route.path == path)
        assert any(dependency.call is developer_only for dependency in route.dependant.dependencies)
    response = download_student_account_template(actor)
    workbook = load_workbook(BytesIO(response.body))
    assert [cell.value for cell in workbook.active[1]][:2] == ["学号*", "姓名*"]
    assert workbook.active["A2"].number_format == "@"


def test_activation_manual_creation_and_reset_use_full_student_number(database):
    db, actor = database
    roster = StudentRoster(student_no="2026999060", name="激活学生", gender="男",
                           college="学院", major="专业", class_name="班级", activated=False)
    db.add(roster)
    db.commit()
    student_register(StudentRegisterRequest(student_no=roster.student_no, name=roster.name), db)
    activated = student(db, roster.student_no)
    assert verify_password("DF2026999060", activated.password_hash)
    activated.password_hash = hash_password("Changed123456!")
    activated.failed_login_count = 4
    db.commit()
    reset_password(activated.id, ResetPasswordRequest(), actor, db)
    db.refresh(activated)
    assert verify_password("DF2026999060", activated.password_hash)
    assert activated.failed_login_count == 0 and activated.locked_until is None
    create_account(AccountCreate(username="2026999061", display_name="手动学生",
                                 password="Ignored123456!", role=Role.STUDENT), actor, db)
    assert verify_password("DF2026999061", student(db, "2026999061").password_hash)


def test_http_upload_authorization_and_student_login(database):
    db, actor = database
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    def override_db():
        yield db
    app.dependency_overrides[get_db] = override_db
    developer_token = _login(db, LoginRequest(username=actor.username, password="Dev123456!"), {Role.DEVELOPER})["access_token"]
    teacher = User(username="teacher", display_name="老师", role=Role.SPORTS_TEACHER,
                   password_hash=actor.password_hash, status="ACTIVE")
    db.add(teacher)
    db.commit()
    teacher_token = _login(db, LoginRequest(username="teacher", password="Dev123456!"), {Role.SPORTS_TEACHER})["access_token"]
    with TestClient(app) as client:
        assert client.get("/api/v1/developer/students/template").status_code == 401
        denied = client.post("/api/v1/developer/students/import",
                             headers={"Authorization": f"Bearer {teacher_token}"},
                             files={"file": ("名单.xlsx", upload([["2026999070", "接口学生"]]).file.read())})
        assert denied.status_code == 403
        result = client.post("/api/v1/developer/students/import",
                             headers={"Authorization": f"Bearer {developer_token}"},
                             files={"file": ("名单.xlsx", upload([["2026999070", "接口学生"]]).file.read())})
        assert result.status_code == 200 and result.json()["registered"] == 1
        login = client.post("/api/v1/auth/student-login", json={"username": "2026999070", "password": "DF2026999070"})
        assert login.status_code == 200 and login.json()["first_login_change_required"] is False
